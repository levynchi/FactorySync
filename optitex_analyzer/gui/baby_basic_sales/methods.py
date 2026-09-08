import os
from datetime import datetime
from tkinter import messagebox, filedialog

from optitex_analyzer.core.data_processor import parse_numeric_value
from .. import theme


class BabyBasicSalesMethodsMixin:
    """Logic for Baby Basic wholesale notes, price list and account."""

    BB_ROOM_COUNT_EXTRA_MODELS = ('אוברול', 'חולצות טורקיה', 'גופייה טורקיה')
    BB_ROOM_COUNT_COLORS = (
        'חום', 'כחול', 'שחור', 'מנטה', 'אפור בהיר', 'לבן',
        'ורוד בהיר', 'כחול נייבי', 'כחול מלאנז', 'ורוד בייבי מלאנז', 'פוקסיה',
    )
    BB_ROOM_COUNT_PRINTS = (
        'משולשים', 'מנומר', 'נקודות', 'חלל', 'באר שבע בנות',
        'פרפרים', 'באר שבע בנים', 'נוצות', 'אוהלים', 'עלים',
    )
    BB_ROOM_COUNT_FABRICS = ('פלנל', 'טריקו', 'ופל')

    def _bb_money(self, value) -> str:
        try:
            return f"{parse_numeric_value(value):,.2f}"
        except Exception:
            return "0.00"

    def _bb_refresh_customer_combos(self):
        partner = self._bb_partner_name()
        names = []
        try:
            names = self.data_processor.get_baby_basic_customers()
        except Exception:
            names = []
        if partner not in names:
            names = [partner] + [n for n in names if n != partner]
        if getattr(self, 'bb_customer_combo', None) is not None:
            try:
                self.bb_customer_combo['values'] = names
                current = (self.bb_customer_var.get() if hasattr(self, 'bb_customer_var') else '').strip()
                if not current:
                    self.bb_customer_var.set(partner)
            except Exception:
                pass

    def _bb_products(self):
        try:
            return self.data_processor.get_baby_basic_products()
        except Exception:
            return []

    def _bb_product_by_barcode(self, barcode: str):
        barcode = str(barcode or '').strip()
        for p in self._bb_products():
            if p.get('barcode') == barcode:
                return p
        return None

    def _bb_filter_products(self, query: str = ''):
        q = (query or '').strip().lower()
        terms = q.split()
        result = []
        for p in self._bb_products():
            blob = ' '.join([
                str(p.get('print_name') or ''),
                str(p.get('item_name') or ''),
                str(p.get('size') or ''),
                str(p.get('fabric') or ''),
                str(p.get('barcode') or ''),
                str(p.get('color') or ''),
            ]).lower()
            if terms and not all(t in blob for t in terms):
                continue
            result.append(p)
        return result

    def _bb_clear_tree(self, tree):
        if tree is None:
            return
        for iid in tree.get_children():
            tree.delete(iid)

    def _bb_inline_edit(self, tree, event, column_id, on_commit):
        col = tree.identify_column(event.x)
        row = tree.identify_row(event.y)
        if not row or col != column_id:
            return
        bbox = tree.bbox(row, col)
        if not bbox:
            return
        x, y, w, h = bbox
        vals = list(tree.item(row, 'values'))
        col_index = int(column_id.replace('#', '')) - 1
        current = vals[col_index] if 0 <= col_index < len(vals) else ''
        import tkinter as tk
        entry = tk.Entry(tree, justify='center')
        entry.insert(0, '' if str(current) in ('0', '0.00') else str(current))
        entry.select_range(0, 'end')
        entry.place(x=x, y=y, width=w, height=h)
        entry.focus_set()

        def commit(_e=None):
            raw = entry.get().strip()
            try:
                entry.destroy()
            except Exception:
                pass
            on_commit(row, raw)

        entry.bind('<Return>', commit)
        entry.bind('<KP_Enter>', commit)
        entry.bind('<FocusOut>', commit)
        entry.bind('<Escape>', lambda _e: entry.destroy())

    # ---------- מחירון ----------
    def _bb_refresh_price_table(self):
        tree = getattr(self, 'bb_price_tree', None)
        if tree is None:
            return
        self._bb_clear_tree(tree)
        query = self.bb_price_search_var.get() if hasattr(self, 'bb_price_search_var') else ''
        for p in self._bb_filter_products(query):
            tree.insert('', 'end', iid=p['barcode'], values=(
                p.get('print_name') or p.get('item_name') or '',
                p.get('size', ''),
                p.get('fabric', ''),
                p.get('color', ''),
                p.get('pack_qty', 5),
                p.get('barcode', ''),
                self._bb_money(p.get('price')),
            ))
        theme.stripe_tree(tree)

    def _bb_edit_price_cell(self, event):
        def on_commit(barcode, raw):
            try:
                price = parse_numeric_value(raw)
            except Exception:
                price = 0
            self.data_processor.set_baby_basic_price(barcode, price)
            self._bb_refresh_price_table()
            self._bb_refresh_note_products()
        self._bb_inline_edit(self.bb_price_tree, event, '#7', on_commit)

    def _bb_apply_price_to_family(self):
        tree = getattr(self, 'bb_price_tree', None)
        if tree is None:
            return
        sel = tree.selection()
        if not sel:
            messagebox.showwarning("מחירון", "בחר מוצר בטבלה ואז החל את המחיר לכל המידות של אותו דגם.")
            return
        src = self._bb_product_by_barcode(sel[0])
        if not src:
            return
        try:
            price = parse_numeric_value(tree.item(sel[0], 'values')[-1])
        except Exception:
            price = src.get('price') or 0
        if price <= 0:
            messagebox.showwarning("מחירון", "קודם הזן מחיר לשורה שנבחרה (דאבל-קליק על עמודת המחיר).")
            return
        family_key = (src.get('print_name') or src.get('item_name') or '').strip()
        fabric = (src.get('fabric') or '').strip()
        updates = {}
        for p in self._bb_products():
            same_name = (p.get('print_name') or p.get('item_name') or '').strip() == family_key
            same_fabric = (p.get('fabric') or '').strip() == fabric
            if same_name and same_fabric:
                updates[p['barcode']] = price
        if not updates:
            return
        self.data_processor.set_baby_basic_prices_bulk(updates)
        self._bb_refresh_price_table()
        self._bb_refresh_note_products()
        messagebox.showinfo("מחירון", f"עודכן מחיר {self._bb_money(price)} ₪ ל-{len(updates)} פריטים של «{family_key}».")

    # ---------- תעודת סחורה ----------
    def _bb_refresh_note_products(self):
        tree = getattr(self, 'bb_note_products_tree', None)
        if tree is None:
            return
        self._bb_clear_tree(tree)
        query = self.bb_note_search_var.get() if hasattr(self, 'bb_note_search_var') else ''
        qty_map = getattr(self, '_bb_note_qty_by_barcode', {})
        for p in self._bb_filter_products(query):
            bc = p['barcode']
            tree.insert('', 'end', iid=bc, values=(
                p.get('print_name') or p.get('item_name') or '',
                p.get('size', ''),
                p.get('fabric', ''),
                bc,
                self._bb_money(p.get('price')),
                int(qty_map.get(bc, 0) or 0),
            ))
        theme.stripe_tree(tree)

    def _bb_edit_note_qty_cell(self, event):
        def on_commit(barcode, raw):
            try:
                qty = int(float(raw)) if raw else 0
            except Exception:
                qty = 0
            if qty < 0:
                qty = 0
            if not hasattr(self, '_bb_note_qty_by_barcode'):
                self._bb_note_qty_by_barcode = {}
            if qty > 0:
                self._bb_note_qty_by_barcode[barcode] = qty
            else:
                self._bb_note_qty_by_barcode.pop(barcode, None)
            if self.bb_note_products_tree.exists(barcode):
                vals = list(self.bb_note_products_tree.item(barcode, 'values'))
                vals[-1] = qty
                self.bb_note_products_tree.item(barcode, values=vals)

        self._bb_inline_edit(self.bb_note_products_tree, event, '#6', on_commit)

    def _bb_refresh_note_lines(self):
        tree = getattr(self, 'bb_note_lines_tree', None)
        if tree is None:
            return
        self._bb_clear_tree(tree)
        lines = getattr(self, '_bb_note_lines', [])
        total = 0.0
        qty_total = 0
        for i, line in enumerate(lines):
            total += parse_numeric_value(line.get('line_total'))
            qty_total += int(line.get('quantity') or 0)
            tree.insert('', 'end', iid=str(i), values=(
                line.get('print_name') or line.get('item_name') or '',
                line.get('size', ''),
                line.get('fabric', ''),
                line.get('barcode', ''),
                int(line.get('quantity') or 0),
                self._bb_money(line.get('unit_price')),
                self._bb_money(line.get('line_total')),
            ))
        theme.stripe_tree(tree)
        if hasattr(self, 'bb_note_total_var'):
            self.bb_note_total_var.set(f"סה״כ יחידות: {qty_total}    סה״כ לתשלום: {self._bb_money(total)} ₪")

    def _bb_add_selected_to_note(self):
        qty_map = {bc: q for bc, q in getattr(self, '_bb_note_qty_by_barcode', {}).items() if int(q or 0) > 0}
        if not qty_map:
            messagebox.showwarning("תעודה", "לא הוזנה כמות. לחץ פעמיים על עמודת «כמות» ואז הוסף לתעודה.")
            return
        if not hasattr(self, '_bb_note_lines'):
            self._bb_note_lines = []
        existing = {str(l.get('barcode')): l for l in self._bb_note_lines}
        missing_price = []
        for bc, qty in qty_map.items():
            prod = self._bb_product_by_barcode(bc)
            if not prod:
                continue
            price = parse_numeric_value(prod.get('price'))
            if price <= 0:
                missing_price.append(prod.get('print_name') or prod.get('item_name') or bc)
            qty = int(qty)
            if bc in existing:
                existing[bc]['quantity'] = int(existing[bc].get('quantity') or 0) + qty
                existing[bc]['unit_price'] = price
                existing[bc]['line_total'] = round(existing[bc]['quantity'] * price, 2)
            else:
                self._bb_note_lines.append({
                    'barcode': bc,
                    'item_name': prod.get('item_name', ''),
                    'print_name': prod.get('print_name') or prod.get('item_name') or '',
                    'size': prod.get('size', ''),
                    'fabric': prod.get('fabric', ''),
                    'pack_qty': prod.get('pack_qty', 5),
                    'quantity': qty,
                    'unit_price': price,
                    'line_total': round(qty * price, 2),
                })
        self._bb_note_qty_by_barcode = {}
        self._bb_refresh_note_products()
        self._bb_refresh_note_lines()
        if missing_price:
            messagebox.showwarning(
                "מחירון חסר",
                "הפריטים הבאים נוספו בלי מחיר. עדכן במחירון או דאבל-קליק על מחיר בשורת התעודה:\n• "
                + "\n• ".join(missing_price[:12])
            )

    def _bb_edit_line_price_cell(self, event):
        def on_commit(iid, raw):
            try:
                idx = int(iid)
            except Exception:
                return
            lines = getattr(self, '_bb_note_lines', [])
            if idx < 0 or idx >= len(lines):
                return
            price = parse_numeric_value(raw)
            qty = int(lines[idx].get('quantity') or 0)
            lines[idx]['unit_price'] = round(price, 2)
            lines[idx]['line_total'] = round(qty * price, 2)
            self._bb_refresh_note_lines()
        self._bb_inline_edit(self.bb_note_lines_tree, event, '#6', on_commit)

    def _bb_remove_selected_note_line(self):
        tree = getattr(self, 'bb_note_lines_tree', None)
        if tree is None:
            return
        sel = tree.selection()
        if not sel:
            return
        try:
            idx = int(sel[0])
            if 0 <= idx < len(self._bb_note_lines):
                self._bb_note_lines.pop(idx)
        except Exception:
            return
        self._bb_refresh_note_lines()

    def _bb_clear_current_note(self):
        self._bb_note_lines = []
        self._bb_note_qty_by_barcode = {}
        if hasattr(self, 'bb_note_comment_var'):
            self.bb_note_comment_var.set('')
        if hasattr(self, 'bb_note_date_var'):
            self.bb_note_date_var.set(datetime.now().strftime('%Y-%m-%d'))
        self._bb_refresh_note_products()
        self._bb_refresh_note_lines()

    def _bb_save_note(self):
        customer = (self.bb_customer_var.get() if hasattr(self, 'bb_customer_var') else '').strip()
        if not customer:
            customer = self._bb_partner_name()
            if hasattr(self, 'bb_customer_var'):
                self.bb_customer_var.set(customer)
        date_str = (self.bb_note_date_var.get() if hasattr(self, 'bb_note_date_var') else '').strip()
        comment = (self.bb_note_comment_var.get() if hasattr(self, 'bb_note_comment_var') else '').strip()
        lines = list(getattr(self, '_bb_note_lines', []) or [])
        try:
            new_id = self.data_processor.add_baby_basic_note(customer, date_str, lines, note=comment)
        except Exception as e:
            messagebox.showerror("שגיאה", str(e))
            return
        self._bb_refresh_customer_combos()
        self._bb_refresh_saved_notes()
        self._bb_refresh_account()
        self._bb_clear_current_note()
        messagebox.showinfo("נשמר", f"תעודת סחורה #{new_id} נשמרה.")
        rec = self.data_processor.get_baby_basic_note(new_id)
        if rec:
            self._bb_open_note_view(rec)

    def _bb_refresh_saved_notes(self):
        tree = getattr(self, 'bb_saved_notes_tree', None)
        if tree is None:
            return
        self._bb_clear_tree(tree)
        notes = sorted(self.data_processor.get_baby_basic_notes(), key=lambda n: int(n.get('id') or 0), reverse=True)
        for rec in notes:
            tree.insert('', 'end', iid=str(rec.get('id')), values=(
                rec.get('id', ''),
                rec.get('date', ''),
                rec.get('customer', ''),
                rec.get('total_quantity', 0),
                self._bb_money(rec.get('total_amount')),
                rec.get('note', ''),
            ))
        theme.stripe_tree(tree)

    def _bb_open_selected_note(self, event=None):
        tree = getattr(self, 'bb_saved_notes_tree', None)
        if tree is None:
            return
        sel = tree.selection()
        if not sel:
            return
        rec = self.data_processor.get_baby_basic_note(sel[0])
        if rec:
            self._bb_open_note_view(rec)

    def _bb_delete_selected_note(self):
        tree = getattr(self, 'bb_saved_notes_tree', None)
        if tree is None:
            return
        sel = tree.selection()
        if not sel:
            return
        rec = self.data_processor.get_baby_basic_note(sel[0])
        if not rec:
            return
        if not messagebox.askyesno("מחיקה", f"למחוק תעודה #{rec.get('id')} לשותף {rec.get('customer')}?"):
            return
        self.data_processor.delete_baby_basic_note(rec.get('id'))
        self._bb_refresh_saved_notes()
        self._bb_refresh_account()
        self._bb_refresh_customer_combos()

    def _bb_open_note_view(self, rec: dict):
        import tkinter as tk
        from tkinter import ttk
        win = tk.Toplevel(self.root)
        win.title(f"תעודת סחורה בייבי בייסיק #{rec.get('id')}")
        win.configure(bg=theme.PAGE_BG)
        win.geometry("920x560")
        tk.Label(
            win,
            text=f"תעודת סחורה #{rec.get('id')}  |  {rec.get('customer')}  |  {rec.get('date')}",
            font=(theme.FONT_FAMILY, 13, 'bold'),
            bg=theme.PAGE_BG,
            fg=theme.DARK,
        ).pack(pady=(10, 4))
        if rec.get('note'):
            tk.Label(win, text=rec.get('note'), bg=theme.PAGE_BG, fg=theme.SUBTEXT).pack()

        cols = ('print_name', 'size', 'fabric', 'color', 'barcode', 'qty', 'price', 'total')
        headers = {
            'print_name': 'מוצר', 'size': 'מידה', 'fabric': 'בד', 'color': 'צבע',
            'barcode': 'ברקוד', 'qty': 'יחידות', 'price': 'מחיר ליחידה', 'total': 'סה״כ',
        }
        tree = theme.make_treeview(win, columns=cols, show='headings', height=14)
        for c in cols:
            tree.heading(c, text=headers[c])
            tree.column(c, width=90 if c in ('size', 'qty', 'price', 'total', 'color') else 160, anchor='center')
        vs = ttk.Scrollbar(win, orient='vertical', command=tree.yview)
        tree.configure(yscrollcommand=vs.set)
        tree.pack(side='left', fill='both', expand=True, padx=(12, 0), pady=8)
        vs.pack(side='left', fill='y', pady=8)
        for line in rec.get('lines') or []:
            tree.insert('', 'end', values=(
                line.get('print_name') or line.get('item_name') or '',
                line.get('size', ''),
                line.get('fabric', ''),
                line.get('color', '') or '',
                line.get('barcode', ''),
                line.get('quantity', 0),
                self._bb_money(line.get('unit_price')),
                self._bb_money(line.get('line_total')),
            ))
        theme.stripe_tree(tree)

        footer = tk.Frame(win, bg=theme.PAGE_BG)
        footer.pack(fill='x', padx=12, pady=(0, 10))
        tk.Label(
            footer,
            text=f"סה״כ יחידות: {rec.get('total_quantity', 0)}    סה״כ: {self._bb_money(rec.get('total_amount'))} ₪",
            font=(theme.FONT_FAMILY, 11, 'bold'),
            bg=theme.PAGE_BG,
        ).pack(side='right')
        theme.make_button(footer, "ייצוא לאקסל", kind="success", command=lambda: self._bb_export_note_excel(rec)).pack(side='left')
        theme.make_button(footer, "PDF להדפסה (A4)", kind="primary", command=lambda: self._bb_export_note_pdf(rec)).pack(side='left', padx=6)

    def _bb_export_note_excel(self, rec: dict, path: str | None = None):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Border, Font, Side, PatternFill
        except Exception as e:
            messagebox.showerror("שגיאה", f"נדרש openpyxl לייצוא לאקסל:\n{e}")
            return
        wb = Workbook()
        ws = wb.active
        ws.title = "תעודת סחורה"
        try:
            ws.sheet_view.rightToLeft = True
        except Exception:
            pass
        ws.page_setup.orientation = 'portrait'
        ws.page_setup.paperSize = 9
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

        thin = Border(
            left=Side(style='thin', color='CBD5E1'),
            right=Side(style='thin', color='CBD5E1'),
            top=Side(style='thin', color='CBD5E1'),
            bottom=Side(style='thin', color='CBD5E1'),
        )
        header_fill = PatternFill('solid', fgColor='1E293B')
        header_font = Font(name='Arial', bold=True, color='FFFFFF', size=11)
        title_font = Font(name='Arial', bold=True, size=16, color='1E293B')
        center = Alignment(horizontal='center', vertical='center')

        biz_name = ''
        try:
            biz_name = (self.settings.get('business.name', '') or '') if getattr(self, 'settings', None) else ''
        except Exception:
            biz_name = ''
        ws.merge_cells('A1:H1')
        ws['A1'] = biz_name or 'בייבי בייסיק'
        ws['A1'].font = title_font
        ws['A1'].alignment = center
        ws.merge_cells('A2:H2')
        ws['A2'] = f"תעודת סחורה #{rec.get('id')}"
        ws['A2'].font = Font(name='Arial', bold=True, size=13)
        ws['A2'].alignment = center
        ws['A4'] = 'לקוח'
        ws['B4'] = rec.get('customer', '')
        ws['C4'] = 'תאריך'
        ws['D4'] = rec.get('date', '')
        if rec.get('note'):
            ws['A5'] = 'הערה'
            ws.merge_cells('B5:H5')
            ws['B5'] = rec.get('note')

        headers = ['מוצר', 'מידה', 'בד', 'צבע', 'ברקוד', 'יחידות', 'מחיר ליחידה', 'סה״כ']
        start_row = 7
        for col, h in enumerate(headers, 1):
            cell = ws.cell(start_row, col, h)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
            cell.border = thin
        for i, line in enumerate(rec.get('lines') or [], 1):
            row = start_row + i
            values = [
                line.get('print_name') or line.get('item_name') or '',
                line.get('size', ''),
                line.get('fabric', ''),
                line.get('color', '') or 'לבן',
                line.get('barcode', ''),
                int(line.get('quantity') or 0),
                round(parse_numeric_value(line.get('unit_price')), 2),
                round(parse_numeric_value(line.get('line_total')), 2),
            ]
            for col, val in enumerate(values, 1):
                cell = ws.cell(row, col, val)
                cell.alignment = center
                cell.border = thin
                if col >= 7:
                    cell.number_format = '#,##0.00'
        total_row = start_row + 1 + len(rec.get('lines') or [])
        ws.cell(total_row, 5, 'סה״כ').font = Font(name='Arial', bold=True)
        ws.cell(total_row, 6, rec.get('total_quantity', 0)).font = Font(name='Arial', bold=True)
        total_cell = ws.cell(total_row, 8, round(parse_numeric_value(rec.get('total_amount')), 2))
        total_cell.font = Font(name='Arial', bold=True)
        total_cell.number_format = '#,##0.00'
        widths = [28, 12, 12, 10, 18, 12, 14, 14]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w

        export_dir = os.path.join(os.getcwd(), 'exports', 'baby_basic_notes')
        os.makedirs(export_dir, exist_ok=True)
        safe_id = rec.get('id')
        safe_date = str(rec.get('date') or '').replace(':', '-')
        default_name = f"baby_basic_note_{safe_id}_{safe_date}.xlsx"
        if not path:
            path = filedialog.asksaveasfilename(
                title="שמירת תעודת סחורה",
                defaultextension=".xlsx",
                initialdir=export_dir,
                initialfile=default_name,
                filetypes=[("Excel", "*.xlsx")],
            )
        if not path:
            return
        wb.save(path)
        try:
            os.startfile(path)
        except Exception:
            pass

    def _bb_print_selected_note_pdf(self):
        tree = getattr(self, 'bb_saved_notes_tree', None)
        if tree is None:
            return
        sel = tree.selection()
        if not sel:
            messagebox.showwarning("הדפסה", "בחר תעודה מהרשימה.")
            return
        rec = self.data_processor.get_baby_basic_note(sel[0])
        if rec:
            self._bb_export_note_pdf(rec)

    def _bb_export_note_pdf(self, rec: dict, path: str | None = None):
        try:
            from optitex_analyzer.core.baby_basic_note_pdf import generate_baby_basic_note_pdf
        except Exception as e:
            messagebox.showerror("שגיאה", f"לא ניתן ליצור PDF:\n{e}")
            return
        export_dir = os.path.join(os.getcwd(), 'exports', 'baby_basic_notes')
        os.makedirs(export_dir, exist_ok=True)
        safe_id = rec.get('id')
        safe_date = str(rec.get('date') or '').replace(':', '-')
        default_name = f"baby_basic_note_{safe_id}_{safe_date}.pdf"
        if not path:
            path = os.path.join(export_dir, default_name)
        biz_name = ''
        try:
            biz_name = (self.settings.get('business.name', '') or '') if getattr(self, 'settings', None) else ''
        except Exception:
            biz_name = ''
        try:
            generate_baby_basic_note_pdf(rec, path, biz_name=biz_name or 'בייבי בייסיק')
        except Exception as e:
            messagebox.showerror("שגיאה", f"יצירת PDF נכשלה:\n{e}")
            return
        try:
            os.startfile(path)
        except Exception:
            messagebox.showinfo("PDF", f"הקובץ נשמר:\n{path}")

    # ---------- חשבון שותף: בייבי בייסיק ----------
    def _bb_partner_name(self) -> str:
        return getattr(self.data_processor, 'BABY_BASIC_PARTNER_NAME', 'בייבי בייסיק')

    def _bb_refresh_account(self):
        acc = self.data_processor.get_baby_basic_account()
        if hasattr(self, 'bb_acc_supplied_var'):
            self.bb_acc_supplied_var.set(f"{self._bb_money(acc['supplied'])} ₪")
        if hasattr(self, 'bb_acc_paid_var'):
            self.bb_acc_paid_var.set(f"{self._bb_money(acc['paid'])} ₪")
        if hasattr(self, 'bb_acc_balance_var'):
            self.bb_acc_balance_var.set(f"{self._bb_money(acc['balance'])} ₪")
        if hasattr(self, 'bb_acc_balance_label'):
            color = theme.DANGER if acc['balance'] > 0.009 else theme.SUCCESS
            self.bb_acc_balance_label.configure(fg=color)

        notes_tree = getattr(self, 'bb_acc_notes_tree', None)
        if notes_tree is not None:
            self._bb_clear_tree(notes_tree)
            notes = sorted(acc['notes'], key=lambda n: (str(n.get('date') or ''), int(n.get('id') or 0)), reverse=True)
            for rec in notes:
                notes_tree.insert('', 'end', iid=f"n{rec.get('id')}", values=(
                    rec.get('date', ''),
                    rec.get('id', ''),
                    rec.get('customer', ''),
                    rec.get('total_quantity', 0),
                    self._bb_money(rec.get('total_amount')),
                    rec.get('note', ''),
                ))
            theme.stripe_tree(notes_tree)

        pay_tree = getattr(self, 'bb_acc_pay_tree', None)
        if pay_tree is not None:
            self._bb_clear_tree(pay_tree)
            pays = sorted(acc['payments'], key=lambda p: (str(p.get('date') or ''), int(p.get('id') or 0)), reverse=True)
            for rec in pays:
                pay_tree.insert('', 'end', iid=str(rec.get('id')), values=(
                    rec.get('date', ''),
                    rec.get('id', ''),
                    rec.get('customer') or self._bb_partner_name(),
                    self._bb_money(rec.get('amount')),
                    rec.get('note', ''),
                ))
            theme.stripe_tree(pay_tree)

        hist_tree = getattr(self, 'bb_acc_hist_tree', None)
        if hist_tree is not None:
            self._bb_clear_tree(hist_tree)
            rows = []
            running = 0.0
            events = []
            for n in acc['notes']:
                events.append(('note', n.get('date'), int(n.get('id') or 0), n))
            for p in acc['payments']:
                events.append(('pay', p.get('date'), int(p.get('id') or 0), p))
            events.sort(key=lambda x: (str(x[1] or ''), 0 if x[0] == 'note' else 1, x[2]))
            for kind, date, _id, rec in events:
                if kind == 'note':
                    amount = parse_numeric_value(rec.get('total_amount'))
                    running += amount
                    rows.append((date, 'תעודת סחורה', rec.get('id'), rec.get('customer') or self._bb_partner_name(), amount, 0, running, rec.get('note', '')))
                else:
                    amount = parse_numeric_value(rec.get('amount'))
                    running -= amount
                    rows.append((date, 'תשלום', rec.get('id'), rec.get('customer') or self._bb_partner_name(), 0, amount, running, rec.get('note', '')))
            for i, row in enumerate(reversed(rows)):
                hist_tree.insert('', 'end', iid=str(i), values=(
                    row[0], row[1], row[2], row[3],
                    self._bb_money(row[4]) if row[4] else '',
                    self._bb_money(row[5]) if row[5] else '',
                    self._bb_money(row[6]),
                    row[7],
                ))
            theme.stripe_tree(hist_tree)

        self._bb_refresh_customer_combos()

    def _bb_add_payment(self):
        date_str = (self.bb_pay_date_var.get() if hasattr(self, 'bb_pay_date_var') else '').strip()
        amount = parse_numeric_value(self.bb_pay_amount_var.get() if hasattr(self, 'bb_pay_amount_var') else 0)
        note = (self.bb_pay_note_var.get() if hasattr(self, 'bb_pay_note_var') else '').strip()
        try:
            self.data_processor.add_baby_basic_payment(date_str=date_str, amount=amount, note=note)
        except Exception as e:
            messagebox.showerror("שגיאה", str(e))
            return
        if hasattr(self, 'bb_pay_amount_var'):
            self.bb_pay_amount_var.set('')
        if hasattr(self, 'bb_pay_note_var'):
            self.bb_pay_note_var.set('')
        self._bb_refresh_account()

    def _bb_delete_selected_payment(self):
        tree = getattr(self, 'bb_acc_pay_tree', None)
        if tree is None:
            return
        sel = tree.selection()
        if not sel:
            return
        if not messagebox.askyesno("מחיקה", f"למחוק תשלום #{sel[0]}?"):
            return
        self.data_processor.delete_baby_basic_payment(sel[0])
        self._bb_refresh_account()

    def _bb_biz_name(self) -> str:
        try:
            return (self.settings.get('business.name', '') or '') if getattr(self, 'settings', None) else ''
        except Exception:
            return ''

    def _bb_print_selected_payment_pdf(self, event=None):
        tree = getattr(self, 'bb_acc_pay_tree', None)
        if tree is None:
            return
        sel = tree.selection()
        if not sel:
            messagebox.showwarning("הדפסה", "בחר תשלום מהרשימה.")
            return
        rec = self.data_processor.get_baby_basic_payment(sel[0])
        if rec:
            self._bb_export_payment_pdf(rec)

    def _bb_export_payment_pdf(self, rec: dict, path: str | None = None):
        try:
            from optitex_analyzer.core.baby_basic_note_pdf import generate_baby_basic_payment_pdf
        except Exception as e:
            messagebox.showerror("שגיאה", f"לא ניתן ליצור PDF:\n{e}")
            return
        export_dir = os.path.join(os.getcwd(), 'exports', 'baby_basic_payments')
        os.makedirs(export_dir, exist_ok=True)
        safe_id = rec.get('id')
        safe_date = str(rec.get('date') or '').replace(':', '-')
        default_name = f"baby_basic_payment_{safe_id}_{safe_date}.pdf"
        if not path:
            path = os.path.join(export_dir, default_name)
        acc = {}
        try:
            acc = self.data_processor.get_baby_basic_account() or {}
        except Exception:
            acc = {}
        try:
            generate_baby_basic_payment_pdf(
                rec,
                path,
                biz_name=self._bb_biz_name() or 'בייבי בייסיק',
                partner=self._bb_partner_name(),
                supplied=acc.get('supplied'),
                paid=acc.get('paid'),
                balance=acc.get('balance'),
            )
        except Exception as e:
            messagebox.showerror("שגיאה", f"יצירת PDF נכשלה:\n{e}")
            return
        try:
            os.startfile(path)
        except Exception:
            messagebox.showinfo("PDF", f"הקובץ נשמר:\n{path}")

    def _bb_print_payments_report_pdf(self):
        try:
            from optitex_analyzer.core.baby_basic_note_pdf import generate_baby_basic_payments_report_pdf
        except Exception as e:
            messagebox.showerror("שגיאה", f"לא ניתן ליצור PDF:\n{e}")
            return
        acc = self.data_processor.get_baby_basic_account() or {}
        payments = acc.get('payments') or self.data_processor.get_baby_basic_payments()
        if not payments:
            messagebox.showwarning("הדפסה", "אין תשלומים להדפסה.")
            return
        export_dir = os.path.join(os.getcwd(), 'exports', 'baby_basic_payments')
        os.makedirs(export_dir, exist_ok=True)
        path = os.path.join(export_dir, f"baby_basic_payments_{datetime.now().strftime('%Y%m%d')}.pdf")
        try:
            generate_baby_basic_payments_report_pdf(
                payments,
                path,
                biz_name=self._bb_biz_name() or 'בייבי בייסיק',
                partner=self._bb_partner_name(),
                supplied=acc.get('supplied'),
                paid=acc.get('paid'),
                balance=acc.get('balance'),
            )
        except Exception as e:
            messagebox.showerror("שגיאה", f"יצירת PDF נכשלה:\n{e}")
            return
        try:
            os.startfile(path)
        except Exception:
            messagebox.showinfo("PDF", f"הקובץ נשמר:\n{path}")

    def _bb_open_account_note(self, event=None):
        tree = getattr(self, 'bb_acc_notes_tree', None)
        if tree is None:
            return
        sel = tree.selection()
        if not sel:
            return
        note_id = str(sel[0]).lstrip('n')
        rec = self.data_processor.get_baby_basic_note(note_id)
        if rec:
            self._bb_open_note_view(rec)

    # ---------- ספירת מלאי חדר (בלי ברקודים) ----------
    def _bb_room_size_sort_key(self, size: str):
        order = ['0-3', '3-6', '6-12', '12-18', '18-24', '24-30']
        s = str(size or '').strip()
        try:
            return (0, order.index(s))
        except ValueError:
            return (1, s)

    def _bb_room_merge_names(self, *groups):
        names = []
        seen = set()
        for group in groups:
            for name in group or []:
                name = str(name or '').strip()
                if name and name not in seen:
                    seen.add(name)
                    names.append(name)
        return names

    def _bb_room_extra_models(self):
        try:
            extra = self.data_processor.get_baby_basic_room_count_extra_models()
        except Exception:
            extra = []
        return self._bb_room_merge_names(self.BB_ROOM_COUNT_EXTRA_MODELS, extra)

    def _bb_room_extra_colors(self):
        try:
            extra = self.data_processor.get_baby_basic_room_count_extra_colors()
        except Exception:
            extra = []
        return self._bb_room_merge_names(self.BB_ROOM_COUNT_COLORS, extra)

    def _bb_room_extra_prints(self):
        try:
            extra = self.data_processor.get_baby_basic_room_count_extra_prints()
        except Exception:
            extra = []
        return self._bb_room_merge_names(self.BB_ROOM_COUNT_PRINTS, extra)

    def _bb_room_extra_fabrics(self):
        try:
            extra = self.data_processor.get_baby_basic_room_count_extra_fabrics()
        except Exception:
            extra = []
        catalog = []
        for p in self._bb_products():
            name = (p.get('fabric') or '').strip()
            if name:
                catalog.append(name)
        return self._bb_room_merge_names(self.BB_ROOM_COUNT_FABRICS, catalog, extra)

    def _bb_room_catalog_for_model(self, model: str = ''):
        model = (model or '').strip()
        sizes = set()
        catalog_models = set()
        extras = self._bb_room_extra_models()
        extra_set = set(extras)
        for p in self._bb_products():
            name = (p.get('print_name') or p.get('item_name') or '').strip()
            if name:
                catalog_models.add(name)
            if model and name != model:
                continue
            size = (p.get('size') or '').strip()
            if size:
                sizes.add(size)
        rest = sorted(m for m in catalog_models if m not in extra_set)
        return {
            'models': extras + rest,
            'colors': self._bb_room_extra_colors(),
            'prints': self._bb_room_extra_prints(),
            'fabrics': self._bb_room_extra_fabrics(),
            'sizes': sorted(sizes, key=self._bb_room_size_sort_key),
        }

    def _bb_refresh_room_combos(self, event=None):
        model = (self.bb_room_model_var.get() if hasattr(self, 'bb_room_model_var') else '').strip()
        opts = self._bb_room_catalog_for_model(model)
        all_opts = self._bb_room_catalog_for_model('')
        if getattr(self, 'bb_room_model_combo', None) is not None:
            try:
                self.bb_room_model_combo['values'] = all_opts['models']
            except Exception:
                pass
        if getattr(self, 'bb_room_color_combo', None) is not None:
            try:
                self.bb_room_color_combo['values'] = all_opts['colors']
            except Exception:
                pass
        if getattr(self, 'bb_room_print_combo', None) is not None:
            try:
                self.bb_room_print_combo['values'] = all_opts['prints']
            except Exception:
                pass
        if getattr(self, 'bb_room_fabric_combo', None) is not None:
            try:
                self.bb_room_fabric_combo['values'] = all_opts['fabrics']
            except Exception:
                pass
        if getattr(self, 'bb_room_size_combo', None) is not None:
            try:
                sizes = opts['sizes'] or all_opts['sizes']
                self.bb_room_size_combo['values'] = sizes
            except Exception:
                pass

    def _bb_on_room_model_change(self, event=None):
        self._bb_refresh_room_combos()

    def _bb_refresh_room_lines(self):
        tree = getattr(self, 'bb_room_lines_tree', None)
        if tree is None:
            return
        self._bb_clear_tree(tree)
        lines = getattr(self, '_bb_room_lines', [])
        total = 0
        for i, line in enumerate(lines):
            qty = int(line.get('quantity') or 0)
            total += qty
            tree.insert('', 'end', iid=str(i), values=(
                line.get('print_name') or '',
                line.get('fabric') or '',
                line.get('color') or '',
                line.get('print') or '',
                line.get('size') or '',
                qty,
            ))
        theme.stripe_tree(tree)
        if hasattr(self, 'bb_room_total_var'):
            self.bb_room_total_var.set(f"סה״כ יחידות: {total}")

    def _bb_add_room_line(self):
        model = (self.bb_room_model_var.get() if hasattr(self, 'bb_room_model_var') else '').strip()
        fabric = (self.bb_room_fabric_var.get() if hasattr(self, 'bb_room_fabric_var') else '').strip()
        color = (self.bb_room_color_var.get() if hasattr(self, 'bb_room_color_var') else '').strip()
        print_val = (self.bb_room_print_var.get() if hasattr(self, 'bb_room_print_var') else '').strip()
        size = (self.bb_room_size_var.get() if hasattr(self, 'bb_room_size_var') else '').strip()
        raw = (self.bb_room_qty_var.get() if hasattr(self, 'bb_room_qty_var') else '').strip()
        if not model:
            messagebox.showwarning("ספירת מלאי", "בחר או הקלד דגם.")
            return
        try:
            qty = int(float(raw)) if raw else 0
        except Exception:
            qty = 0
        if qty <= 0:
            messagebox.showwarning("ספירת מלאי", "הזן כמות גדולה מאפס.")
            return
        if not hasattr(self, '_bb_room_lines'):
            self._bb_room_lines = []
        key = (model, fabric, color, print_val, size)
        merged = False
        for line in self._bb_room_lines:
            existing = (
                str(line.get('print_name') or '').strip(),
                str(line.get('fabric') or '').strip(),
                str(line.get('color') or '').strip(),
                str(line.get('print') or '').strip(),
                str(line.get('size') or '').strip(),
            )
            if existing == key:
                line['quantity'] = int(line.get('quantity') or 0) + qty
                merged = True
                break
        if not merged:
            self._bb_room_lines.append({
                'print_name': model,
                'fabric': fabric,
                'color': color,
                'print': print_val,
                'size': size,
                'quantity': qty,
            })
        self._bb_persist_room_line_extras([{
            'print_name': model,
            'fabric': fabric,
            'color': color,
            'print': print_val,
        }])
        if hasattr(self, 'bb_room_qty_var'):
            self.bb_room_qty_var.set('')
        self._bb_refresh_room_combos()
        self._bb_refresh_room_lines()

    def _bb_edit_room_qty_cell(self, event):
        def on_commit(iid, raw):
            try:
                idx = int(iid)
            except Exception:
                return
            lines = getattr(self, '_bb_room_lines', [])
            if idx < 0 or idx >= len(lines):
                return
            try:
                qty = int(float(raw)) if raw else 0
            except Exception:
                qty = 0
            if qty <= 0:
                lines.pop(idx)
            else:
                lines[idx]['quantity'] = qty
            self._bb_refresh_room_lines()
        self._bb_inline_edit(self.bb_room_lines_tree, event, '#6', on_commit)

    def _bb_remove_selected_room_line(self):
        tree = getattr(self, 'bb_room_lines_tree', None)
        if tree is None:
            return
        sel = tree.selection()
        if not sel:
            return
        try:
            idx = int(sel[0])
            if 0 <= idx < len(self._bb_room_lines):
                self._bb_room_lines.pop(idx)
        except Exception:
            return
        self._bb_refresh_room_lines()

    def _bb_clear_current_room_count(self):
        self._bb_room_lines = []
        if hasattr(self, 'bb_room_note_var'):
            self.bb_room_note_var.set('')
        if hasattr(self, 'bb_room_date_var'):
            self.bb_room_date_var.set(datetime.now().strftime('%Y-%m-%d'))
        if hasattr(self, 'bb_room_qty_var'):
            self.bb_room_qty_var.set('')
        self._bb_refresh_room_lines()

    def _bb_save_room_count(self):
        date_str = (self.bb_room_date_var.get() if hasattr(self, 'bb_room_date_var') else '').strip()
        note = (self.bb_room_note_var.get() if hasattr(self, 'bb_room_note_var') else '').strip()
        lines = list(getattr(self, '_bb_room_lines', []) or [])
        try:
            new_id = self.data_processor.add_baby_basic_room_count(date_str, lines, note=note)
        except Exception as e:
            messagebox.showerror("שגיאה", str(e))
            return
        self._bb_clear_current_room_count()
        self._bb_refresh_saved_room_counts()
        messagebox.showinfo("נשמר", f"ספירת מלאי חדר #{new_id} נשמרה.")
        rec = self.data_processor.get_baby_basic_room_count(new_id)
        if rec:
            self._bb_open_room_count_view(rec)

    def _bb_persist_room_line_extras(self, lines):
        catalog_names = set()
        catalog_fabrics = set(self.BB_ROOM_COUNT_FABRICS)
        for p in self._bb_products():
            name = (p.get('print_name') or p.get('item_name') or '').strip()
            if name:
                catalog_names.add(name)
            fname = (p.get('fabric') or '').strip()
            if fname:
                catalog_fabrics.add(fname)
        for line in lines or []:
            model = str(line.get('print_name') or '').strip()
            fabric = str(line.get('fabric') or '').strip()
            color = str(line.get('color') or '').strip()
            print_val = str(line.get('print') or '').strip()
            if model and model not in catalog_names and model not in self.BB_ROOM_COUNT_EXTRA_MODELS:
                try:
                    self.data_processor.add_baby_basic_room_count_extra_model(model)
                    catalog_names.add(model)
                except Exception:
                    pass
            if fabric and fabric not in catalog_fabrics:
                try:
                    self.data_processor.add_baby_basic_room_count_extra_fabric(fabric)
                    catalog_fabrics.add(fabric)
                except Exception:
                    pass
            if color and color not in self.BB_ROOM_COUNT_COLORS:
                try:
                    self.data_processor.add_baby_basic_room_count_extra_color(color)
                except Exception:
                    pass
            if print_val and print_val not in self.BB_ROOM_COUNT_PRINTS:
                try:
                    self.data_processor.add_baby_basic_room_count_extra_print(print_val)
                except Exception:
                    pass

    def _bb_import_room_count_excel(self):
        initial = os.path.join(os.getcwd(), 'exports', 'baby_basic_room_counts')
        if not os.path.isdir(initial):
            initial = os.getcwd()
        path = filedialog.askopenfilename(
            title='ייבוא ספירת מלאי מאקסל',
            initialdir=initial,
            filetypes=[('Excel', '*.xlsx'), ('All files', '*.*')],
        )
        if not path:
            return
        try:
            parsed = self.data_processor.parse_baby_basic_room_count_excel(path)
        except Exception as e:
            messagebox.showerror("שגיאה", f"קריאת הקובץ נכשלה:\n{e}")
            return
        lines = list(parsed.get('lines') or [])
        self._bb_persist_room_line_extras(lines)
        try:
            new_id = self.data_processor.add_baby_basic_room_count(
                parsed.get('date') or '',
                lines,
                note=parsed.get('note') or '',
            )
        except Exception as e:
            messagebox.showerror("שגיאה", str(e))
            return
        self._bb_refresh_room_combos()
        self._bb_refresh_saved_room_counts()
        messagebox.showinfo("יובא", f"ספירת מלאי חדר #{new_id} נשמרה מתוך הקובץ.\n{len(lines)} שורות.")
        rec = self.data_processor.get_baby_basic_room_count(new_id)
        if rec:
            self._bb_open_room_count_view(rec)

    def _bb_refresh_saved_room_counts(self):
        tree = getattr(self, 'bb_saved_room_tree', None)
        if tree is None:
            return
        self._bb_clear_tree(tree)
        counts = sorted(
            self.data_processor.get_baby_basic_room_counts(),
            key=lambda n: int(n.get('id') or 0),
            reverse=True,
        )
        for rec in counts:
            tree.insert('', 'end', iid=str(rec.get('id')), values=(
                rec.get('id', ''),
                rec.get('date', ''),
                rec.get('total_quantity', 0),
                rec.get('note', ''),
            ))
        theme.stripe_tree(tree)

    def _bb_open_selected_room_count(self, event=None):
        tree = getattr(self, 'bb_saved_room_tree', None)
        if tree is None:
            return
        sel = tree.selection()
        if not sel:
            return
        rec = self.data_processor.get_baby_basic_room_count(sel[0])
        if rec:
            self._bb_open_room_count_view(rec)

    def _bb_delete_selected_room_count(self):
        tree = getattr(self, 'bb_saved_room_tree', None)
        if tree is None:
            return
        sel = tree.selection()
        if not sel:
            return
        rec = self.data_processor.get_baby_basic_room_count(sel[0])
        if not rec:
            return
        if not messagebox.askyesno("מחיקה", f"למחוק ספירת מלאי #{rec.get('id')} מתאריך {rec.get('date')}?"):
            return
        self.data_processor.delete_baby_basic_room_count(rec.get('id'))
        self._bb_refresh_saved_room_counts()

    def _bb_previous_room_count(self, rec: dict):
        counts = sorted(
            self.data_processor.get_baby_basic_room_counts(),
            key=lambda n: (str(n.get('date') or ''), int(n.get('id') or 0)),
        )
        prev = None
        rec_id = int(rec.get('id') or 0)
        for item in counts:
            if int(item.get('id') or 0) == rec_id:
                break
            prev = item
        return prev

    def _bb_open_room_count_view(self, rec: dict):
        import tkinter as tk
        from tkinter import ttk
        win = tk.Toplevel(self.root)
        win.title(f"ספירת מלאי חדר #{rec.get('id')}")
        win.configure(bg=theme.PAGE_BG)
        win.geometry("1140x520")
        tk.Label(
            win,
            text=f"ספירת מלאי חדר #{rec.get('id')}  |  {rec.get('date')}",
            font=(theme.FONT_FAMILY, 13, 'bold'),
            bg=theme.PAGE_BG,
            fg=theme.DARK,
        ).pack(pady=(10, 4))
        if rec.get('note'):
            tk.Label(win, text=rec.get('note'), bg=theme.PAGE_BG, fg=theme.SUBTEXT).pack()

        prev = self._bb_previous_room_count(rec)
        prev_map = {}
        if prev:
            for line in prev.get('lines') or []:
                key = (
                    str(line.get('print_name') or '').strip(),
                    str(line.get('fabric') or '').strip(),
                    str(line.get('color') or '').strip(),
                    str(line.get('print') or '').strip(),
                    str(line.get('size') or '').strip(),
                    str(line.get('area') or '').strip(),
                    str(line.get('box') or '').strip(),
                )
                prev_map[key] = int(line.get('quantity') or 0)
            tk.Label(
                win,
                text=f"השוואה לספירה #{prev.get('id')} מתאריך {prev.get('date')}",
                bg=theme.PAGE_BG,
                fg=theme.SUBTEXT,
                font=theme.FONT_SMALL,
            ).pack()

        cols = ('print_name', 'fabric', 'color', 'print', 'size', 'area', 'box', 'qty', 'delta')
        headers = {
            'print_name': 'דגם',
            'fabric': 'סוג בד',
            'color': 'צבע רקע',
            'print': 'הדפס',
            'size': 'מידה',
            'area': 'אזור',
            'box': 'קופסא',
            'qty': 'כמות',
            'delta': 'שינוי',
        }
        tree = theme.make_treeview(win, columns=cols, show='headings', height=14)
        for c in cols:
            tree.heading(c, text=headers[c])
            w = 80 if c in ('size', 'qty', 'delta', 'area', 'box') else (110 if c in ('fabric', 'color', 'print') else 160)
            tree.column(c, width=w, anchor='center')
        vs = ttk.Scrollbar(win, orient='vertical', command=tree.yview)
        tree.configure(yscrollcommand=vs.set)
        tree.pack(side='left', fill='both', expand=True, padx=(12, 0), pady=8)
        vs.pack(side='left', fill='y', pady=8)
        for line in rec.get('lines') or []:
            qty = int(line.get('quantity') or 0)
            key = (
                str(line.get('print_name') or '').strip(),
                str(line.get('fabric') or '').strip(),
                str(line.get('color') or '').strip(),
                str(line.get('print') or '').strip(),
                str(line.get('size') or '').strip(),
                str(line.get('area') or '').strip(),
                str(line.get('box') or '').strip(),
            )
            delta = ''
            if prev_map:
                old = prev_map.get(key, 0)
                diff = qty - old
                if diff > 0:
                    delta = f"+{diff}"
                elif diff < 0:
                    delta = str(diff)
                else:
                    delta = '0'
            tree.insert('', 'end', values=(
                line.get('print_name') or '',
                line.get('fabric') or '',
                line.get('color') or '',
                line.get('print') or '',
                line.get('size') or '',
                line.get('area') or '',
                line.get('box') or '',
                qty,
                delta,
            ))
        theme.stripe_tree(tree)

        footer = tk.Frame(win, bg=theme.PAGE_BG)
        footer.pack(fill='x', padx=12, pady=(0, 10))
        tk.Label(
            footer,
            text=f"סה״כ יחידות: {rec.get('total_quantity', 0)}",
            font=(theme.FONT_FAMILY, 11, 'bold'),
            bg=theme.PAGE_BG,
        ).pack(side='right')
        theme.make_button(footer, "PDF להדפסה (A4)", kind="primary", command=lambda: self._bb_export_room_count_pdf(rec)).pack(side='left')

    def _bb_print_selected_room_count_pdf(self):
        tree = getattr(self, 'bb_saved_room_tree', None)
        if tree is None:
            return
        sel = tree.selection()
        if not sel:
            messagebox.showwarning("הדפסה", "בחר ספירה מהרשימה.")
            return
        rec = self.data_processor.get_baby_basic_room_count(sel[0])
        if rec:
            self._bb_export_room_count_pdf(rec)

    def _bb_export_room_count_pdf(self, rec: dict, path: str | None = None):
        try:
            from optitex_analyzer.core.baby_basic_note_pdf import generate_baby_basic_room_count_pdf
        except Exception as e:
            messagebox.showerror("שגיאה", f"לא ניתן ליצור PDF:\n{e}")
            return
        export_dir = os.path.join(os.getcwd(), 'exports', 'baby_basic_room_counts')
        os.makedirs(export_dir, exist_ok=True)
        safe_id = rec.get('id')
        safe_date = str(rec.get('date') or '').replace(':', '-')
        default_name = f"baby_basic_room_count_{safe_id}_{safe_date}.pdf"
        if not path:
            path = os.path.join(export_dir, default_name)
        try:
            generate_baby_basic_room_count_pdf(rec, path, biz_name=self._bb_biz_name() or 'בייבי בייסיק')
        except Exception as e:
            messagebox.showerror("שגיאה", f"יצירת PDF נכשלה:\n{e}")
            return
        try:
            os.startfile(path)
        except Exception:
            messagebox.showinfo("PDF", f"הקובץ נשמר:\n{path}")

    def _bb_print_count_sheet_pdf(self):
        try:
            from optitex_analyzer.core.baby_basic_note_pdf import generate_baby_basic_count_sheet_pdf
        except Exception as e:
            messagebox.showerror("שגיאה", f"לא ניתן ליצור PDF:\n{e}")
            return
        export_dir = os.path.join(os.getcwd(), 'exports', 'baby_basic_room_counts')
        os.makedirs(export_dir, exist_ok=True)
        path = os.path.join(
            export_dir,
            f"baby_basic_count_sheet_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf",
        )
        try:
            generate_baby_basic_count_sheet_pdf(
                file_path=path,
                biz_name=self._bb_biz_name() or 'בייבי בייסיק',
                location='חדר בייבי בייסיק',
            )
        except Exception as e:
            messagebox.showerror("שגיאה", f"יצירת PDF נכשלה:\n{e}")
            return
        try:
            os.startfile(path)
        except Exception:
            messagebox.showinfo("PDF", f"הקובץ נשמר:\n{path}")
