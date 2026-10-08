import os
import tkinter as tk
from datetime import datetime
from tkinter import ttk, messagebox, filedialog
from . import theme


class SuppliersTabMixin:
    """Mixin לטאב ספקים: הוספה, מחיקה, הצגה ומסמכים."""

    def _create_suppliers_tab(self):
        tab = tk.Frame(self.notebook, bg=theme.PAGE_BG)
        self.notebook.add(tab, text="ספקים")
        tk.Label(tab, text="ניהול ספקים", font=(theme.FONT_FAMILY,16,'bold'), bg=theme.PAGE_BG, fg=theme.DARK).pack(pady=8)

        form = ttk.LabelFrame(tab, text="הוספת ספק", padding=10)
        form.pack(fill='x', padx=10, pady=6)

        self.sup_business_name_var = tk.StringVar()
        self.sup_first_name_var = tk.StringVar()
        self.sup_phone_var = tk.StringVar()
        self.sup_address_var = tk.StringVar()
        self.sup_business_var = tk.StringVar()
        self.sup_notes_var = tk.StringVar()

        labels = [
            ("שם עסק", self.sup_business_name_var, 18),
            ("שם פרטי", self.sup_first_name_var, 14),
            ("טלפון", self.sup_phone_var, 14),
            ("כתובת", self.sup_address_var, 25),
            ("מס' עסק", self.sup_business_var, 14),
            ("הערות", self.sup_notes_var, 25),
        ]
        for i,(lbl,var,w) in enumerate(labels):
            tk.Label(form, text=f"{lbl}:", font=(theme.FONT_FAMILY,10,'bold')).grid(row=0,column=i*2,sticky='w',padx=4,pady=4)
            tk.Entry(form, textvariable=var, width=w).grid(row=0,column=i*2+1,sticky='w',padx=2,pady=4)

        tk.Button(form, text="➕ הוסף", command=self._add_supplier_record, bg=theme.SUCCESS, fg='white').grid(row=0, column=len(labels)*2, padx=8)
        tk.Button(form, text="📄 מסמכים", command=self._open_selected_supplier_documents, bg=theme.PRIMARY, fg='white').grid(row=0, column=len(labels)*2+1, padx=4)
        tk.Button(form, text="💵 מאזן", command=self._open_selected_supplier_balance, bg=theme.TEAL, fg='white').grid(row=0, column=len(labels)*2+2, padx=4)
        tk.Button(form, text="🗑️ מחק נבחר", command=self._delete_selected_supplier, bg=theme.WARNING, fg='white').grid(row=0, column=len(labels)*2+3, padx=4)

        tree_frame = ttk.LabelFrame(tab, text="ספקים", padding=6)
        tree_frame.pack(fill='both', expand=True, padx=10, pady=6)
        cols = ('id','business_name','first_name','phone','address','business_number','notes','files','balance','created_at')
        self.suppliers_tree = ttk.Treeview(tree_frame, columns=cols, show='headings', height=14)
        headers = {
            'id':'ID','business_name':'שם עסק','first_name':'שם פרטי','phone':'טלפון','address':'כתובת',
            'business_number':'מספר עסק','notes':'הערות','files':'קבצים','balance':'מאזן $','created_at':'נוצר'
        }
        widths = {'id':45,'business_name':150,'first_name':110,'phone':110,'address':200,'business_number':110,'notes':160,'files':60,'balance':100,'created_at':140}
        for c in cols:
            self.suppliers_tree.heading(c, text=headers[c])
            self.suppliers_tree.column(c, width=widths[c], anchor='center')
        vs = ttk.Scrollbar(tree_frame, orient='vertical', command=self.suppliers_tree.yview)
        self.suppliers_tree.configure(yscroll=vs.set)
        self.suppliers_tree.pack(side='left', fill='both', expand=True)
        vs.pack(side='right', fill='y')
        self.suppliers_tree.bind('<Double-1>', lambda _e: self._open_selected_supplier_documents())
        self._load_suppliers_into_tree()

    def _supplier_files_count(self, rec):
        docs = rec.get('documents') if isinstance(rec, dict) else None
        return len(docs) if isinstance(docs, list) else 0

    @staticmethod
    def _supplier_money(value) -> str:
        try:
            return f"${float(value or 0):,.2f}"
        except (TypeError, ValueError):
            return "$0.00"

    def _supplier_balance_text(self, rec):
        """מאזן לתצוגה בטבלה — ריק אם אין רשומות בכרטסת."""
        try:
            if not self.data_processor.get_supplier_ledger(rec.get('id')):
                return ''
            bal = self.data_processor.get_supplier_balance(rec.get('id'))
            return self._supplier_money(bal.get('balance_net'))
        except Exception:
            return ''

    def _load_suppliers_into_tree(self):
        if not hasattr(self, 'suppliers_tree'): return
        for item in self.suppliers_tree.get_children():
            self.suppliers_tree.delete(item)
        try:
            for rec in getattr(self.data_processor, 'suppliers', []):
                self.suppliers_tree.insert('', 'end', values=(
                    rec.get('id'),
                    rec.get('business_name') or rec.get('name'),
                    rec.get('first_name',''),
                    rec.get('phone'),
                    rec.get('address'),
                    rec.get('business_number'),
                    rec.get('notes'),
                    self._supplier_files_count(rec),
                    self._supplier_balance_text(rec),
                    rec.get('created_at')
                ))
        except Exception:
            pass

    def _add_supplier_record(self):
        business_name = self.sup_business_name_var.get().strip()
        if not business_name:
            messagebox.showerror("שגיאה", "חובה להזין שם עסק")
            return
        try:
            self.data_processor.add_supplier(
                business_name,
                self.sup_phone_var.get(),
                self.sup_address_var.get(),
                self.sup_business_var.get(),
                self.sup_notes_var.get(),
                self.sup_first_name_var.get(),
            )
            self._load_suppliers_into_tree()
            self.sup_business_name_var.set(''); self.sup_first_name_var.set(''); self.sup_phone_var.set(''); self.sup_address_var.set(''); self.sup_business_var.set(''); self.sup_notes_var.set('')
            # עדכן קומבואים בטאבים אחרים
            if hasattr(self, '_notify_suppliers_changed'):
                try: self._notify_suppliers_changed()
                except Exception: pass
        except Exception as e:
            messagebox.showerror("שגיאה", str(e))

    def _delete_selected_supplier(self):
        sel = self.suppliers_tree.selection()
        if not sel:
            return
        ids = []
        for item in sel:
            vals = self.suppliers_tree.item(item, 'values')
            if vals:
                ids.append(int(vals[0]))
        deleted_any = False
        for _id in ids:
            if self.data_processor.delete_supplier(_id):
                deleted_any = True
        if deleted_any:
            self._load_suppliers_into_tree()
            if hasattr(self, '_notify_suppliers_changed'):
                try: self._notify_suppliers_changed()
                except Exception: pass

    def _selected_supplier_id(self):
        sel = self.suppliers_tree.selection()
        if not sel:
            return None
        vals = self.suppliers_tree.item(sel[0], 'values')
        if not vals:
            return None
        try:
            return int(vals[0])
        except (TypeError, ValueError):
            return None

    def _open_selected_supplier_documents(self):
        supplier_id = self._selected_supplier_id()
        if supplier_id is None:
            messagebox.showinfo("מסמכים", "יש לבחור ספק מהטבלה")
            return
        supplier = self.data_processor.get_supplier(supplier_id)
        if not supplier:
            messagebox.showerror("שגיאה", "ספק לא נמצא")
            return
        self._show_supplier_documents_dialog(supplier)

    def _open_selected_supplier_balance(self):
        supplier_id = self._selected_supplier_id()
        if supplier_id is None:
            messagebox.showinfo("מאזן", "יש לבחור ספק מהטבלה")
            return
        supplier = self.data_processor.get_supplier(supplier_id)
        if not supplier:
            messagebox.showerror("שגיאה", "ספק לא נמצא")
            return
        self._show_supplier_balance_dialog(supplier)

    def _show_supplier_documents_dialog(self, supplier):
        supplier_id = int(supplier.get('id'))
        display_name = supplier.get('business_name') or supplier.get('name') or f"ספק {supplier_id}"

        dialog = tk.Toplevel()
        dialog.title(f"מסמכי ספק — {display_name}")
        dialog.geometry("760x460")
        dialog.minsize(640, 360)
        dialog.transient(self.root if hasattr(self, 'root') else None)
        dialog.grab_set()
        dialog.configure(bg=theme.PAGE_BG)

        main = tk.Frame(dialog, bg=theme.PAGE_BG)
        main.pack(fill='both', expand=True, padx=16, pady=16)

        tk.Label(
            main,
            text=f"מסמכים של {display_name}",
            font=(theme.FONT_FAMILY, 13, 'bold'),
            bg=theme.PAGE_BG,
            fg=theme.DARK,
        ).pack(anchor='w', pady=(0, 4))
        tk.Label(
            main,
            text="אפשר לייצא שיחה מוואטסאפ (ייצוא צ'אט, ללא מדיה) ולהעלות כאן קובץ txt או zip.",
            font=(theme.FONT_FAMILY, 9),
            bg=theme.PAGE_BG,
            fg=theme.SUBTEXT,
            wraplength=580,
            justify='right',
        ).pack(anchor='w', pady=(0, 10))

        tree_wrap = tk.Frame(main, bg=theme.PAGE_BG)
        tree_wrap.pack(fill='both', expand=True)
        cols = ('original_name', 'uploaded_at')
        docs_tree = ttk.Treeview(tree_wrap, columns=cols, show='headings', height=10, selectmode='extended')
        docs_tree.heading('original_name', text='שם קובץ')
        docs_tree.heading('uploaded_at', text='תאריך העלאה')
        docs_tree.column('original_name', width=380, anchor='w')
        docs_tree.column('uploaded_at', width=160, anchor='center')
        vs = ttk.Scrollbar(tree_wrap, orient='vertical', command=docs_tree.yview)
        docs_tree.configure(yscroll=vs.set)
        docs_tree.pack(side='left', fill='both', expand=True)
        vs.pack(side='right', fill='y')

        def refresh_docs():
            for item in docs_tree.get_children():
                docs_tree.delete(item)
            current = self.data_processor.get_supplier(supplier_id) or {}
            for doc in current.get('documents') or []:
                docs_tree.insert('', 'end', iid=str(doc.get('id')), values=(
                    doc.get('original_name') or doc.get('filename') or '',
                    doc.get('uploaded_at') or '',
                ))
            self._load_suppliers_into_tree()

        def selected_doc_id():
            sel = docs_tree.selection()
            if not sel:
                return None
            try:
                return int(sel[0])
            except (TypeError, ValueError):
                return None

        def find_doc(doc_id):
            current = self.data_processor.get_supplier(supplier_id) or {}
            for doc in current.get('documents') or []:
                try:
                    if int(doc.get('id', 0)) == int(doc_id):
                        return doc
                except (TypeError, ValueError):
                    continue
            return None

        def upload_files():
            paths = filedialog.askopenfilenames(
                title="בחר קבצים לספק",
                filetypes=[
                    ("Excel", "*.xlsx *.xls"),
                    ("WhatsApp export", "*.txt *.zip"),
                    ("PDF", "*.pdf"),
                    ("Word", "*.doc *.docx"),
                    ("תמונות", "*.png *.jpg *.jpeg *.webp *.gif"),
                    ("כל הקבצים", "*.*"),
                ],
            )
            if not paths:
                return
            added = 0
            errors = []
            for path in paths:
                try:
                    self.data_processor.add_supplier_document(supplier_id, path)
                    added += 1
                except Exception as e:
                    errors.append(f"{os.path.basename(path)}: {e}")
            refresh_docs()
            if errors:
                messagebox.showerror("שגיאה", "\n".join(errors), parent=dialog)
            elif added:
                messagebox.showinfo("מסמכים", f"הועלו {added} קבצים", parent=dialog)

        def selected_doc_ids():
            ids = []
            for iid in docs_tree.selection():
                try:
                    ids.append(int(iid))
                except (TypeError, ValueError):
                    continue
            return ids

        def absolute_doc_path(doc):
            path = self.data_processor.supplier_document_path(supplier_id, doc.get('filename') or '')
            return os.path.abspath(path)

        def open_selected():
            doc_id = selected_doc_id()
            if doc_id is None:
                messagebox.showinfo("מסמכים", "יש לבחור קובץ", parent=dialog)
                return
            doc = find_doc(doc_id)
            if not doc:
                messagebox.showerror("שגיאה", "מסמך לא נמצא", parent=dialog)
                return
            self._open_supplier_document_file(absolute_doc_path(doc))

        copy_status = tk.StringVar(value="")

        def copy_selected_paths(_event=None):
            ids = selected_doc_ids()
            if not ids:
                messagebox.showinfo("מסמכים", "יש לבחור קובץ", parent=dialog)
                return "break"
            paths = []
            for doc_id in ids:
                doc = find_doc(doc_id)
                if doc:
                    paths.append(absolute_doc_path(doc))
            if not paths:
                messagebox.showerror("שגיאה", "מסמך לא נמצא", parent=dialog)
                return "break"
            dialog.clipboard_clear()
            dialog.clipboard_append("\n".join(paths))
            dialog.update()
            if len(paths) == 1:
                copy_status.set("הנתיב הועתק ללוח")
            else:
                copy_status.set(f"הועתקו {len(paths)} נתיבים ללוח")
            dialog.after(2500, lambda: copy_status.set(""))
            return "break"

        def show_docs_menu(event):
            row = docs_tree.identify_row(event.y)
            if not row:
                return
            if row not in docs_tree.selection():
                docs_tree.selection_set(row)
                docs_tree.focus(row)
            menu.tk_popup(event.x_root, event.y_root)

        def delete_selected():
            doc_id = selected_doc_id()
            if doc_id is None:
                messagebox.showinfo("מסמכים", "יש לבחור קובץ", parent=dialog)
                return
            doc = find_doc(doc_id)
            label = (doc or {}).get('original_name') or (doc or {}).get('filename') or 'הקובץ'
            if not messagebox.askyesno("מחיקה", f"למחוק את {label}?", parent=dialog):
                return
            if self.data_processor.delete_supplier_document(supplier_id, doc_id):
                refresh_docs()
            else:
                messagebox.showerror("שגיאה", "לא ניתן למחוק את הקובץ", parent=dialog)

        docs_tree.bind('<Double-1>', lambda _e: open_selected())
        docs_tree.bind('<Control-c>', copy_selected_paths)
        docs_tree.bind('<Control-C>', copy_selected_paths)
        docs_tree.bind('<Button-3>', show_docs_menu)

        menu = tk.Menu(dialog, tearoff=0)
        menu.add_command(label="העתק נתיב", command=copy_selected_paths)
        menu.add_command(label="פתח", command=open_selected)
        menu.add_separator()
        menu.add_command(label="מחק קובץ", command=delete_selected)

        buttons = tk.Frame(main, bg=theme.PAGE_BG)
        buttons.pack(fill='x', pady=(12, 0))
        tk.Button(buttons, text="העלה קובץ", command=upload_files, bg=theme.SUCCESS, fg='white', font=(theme.FONT_FAMILY, 10, 'bold'), width=14).pack(side='right', padx=4)
        tk.Button(buttons, text="פתח", command=open_selected, bg=theme.PRIMARY, fg='white', font=(theme.FONT_FAMILY, 10, 'bold'), width=12).pack(side='right', padx=4)
        tk.Button(buttons, text="העתק נתיב", command=copy_selected_paths, bg=theme.TEAL, fg='white', font=(theme.FONT_FAMILY, 10, 'bold'), width=14).pack(side='right', padx=4)
        tk.Button(buttons, text="מחק קובץ", command=delete_selected, bg=theme.DANGER, fg='white', font=(theme.FONT_FAMILY, 10, 'bold'), width=12).pack(side='right', padx=4)
        tk.Button(buttons, text="מאזן", command=lambda: self._show_supplier_balance_dialog(supplier, parent=dialog), bg=theme.DARK_2, fg='white', font=(theme.FONT_FAMILY, 10, 'bold'), width=10).pack(side='right', padx=4)
        tk.Button(buttons, text="סגור", command=dialog.destroy, bg=theme.DARK_2, fg='white', font=(theme.FONT_FAMILY, 10), width=10).pack(side='left', padx=4)
        tk.Label(buttons, textvariable=copy_status, bg=theme.PAGE_BG, fg=theme.TEAL, font=(theme.FONT_FAMILY, 9)).pack(side='left', padx=8)

        refresh_docs()

    def _open_supplier_document_file(self, file_path):
        try:
            if os.path.exists(file_path):
                os.startfile(file_path)
            else:
                messagebox.showerror("שגיאה", f"קובץ לא נמצא: {file_path}")
        except Exception as e:
            messagebox.showerror("שגיאה", f"לא ניתן לפתוח את הקובץ: {str(e)}")

    # ------------------------------------------------------------------
    # מאזן חיובים / תשלומים לספק
    # ------------------------------------------------------------------
    def _show_supplier_balance_dialog(self, supplier, parent=None):
        supplier_id = int(supplier.get('id'))
        display_name = supplier.get('business_name') or supplier.get('name') or f"ספק {supplier_id}"
        money = self._supplier_money

        dialog = tk.Toplevel()
        dialog.title(f"מאזן ספק — {display_name}")
        dialog.geometry("980x620")
        dialog.minsize(820, 480)
        dialog.transient(parent or (self.root if hasattr(self, 'root') else None))
        dialog.grab_set()
        dialog.configure(bg=theme.PAGE_BG)

        def on_close():
            dialog.destroy()
            if parent is not None:
                try:
                    parent.grab_set()
                except Exception:
                    pass
            self._load_suppliers_into_tree()
        dialog.protocol("WM_DELETE_WINDOW", on_close)

        main = tk.Frame(dialog, bg=theme.PAGE_BG)
        main.pack(fill='both', expand=True, padx=16, pady=12)

        head = tk.Frame(main, bg=theme.PAGE_BG)
        head.pack(fill='x')
        tk.Label(head, text=f"מאזן חיובים ותשלומים — {display_name}", font=(theme.FONT_FAMILY, 13, 'bold'), bg=theme.PAGE_BG, fg=theme.DARK).pack(side='right')
        tk.Label(head, text="סכומים בדולרים.", font=theme.FONT_SMALL, bg=theme.PAGE_BG, fg=theme.SUBTEXT).pack(side='right', padx=10)

        cards = tk.Frame(main, bg=theme.PAGE_BG)
        cards.pack(fill='x', pady=(8, 2))

        def make_card(title, var, color):
            box = tk.Frame(cards, bg=theme.CARD_BG, highlightbackground=theme.BORDER, highlightthickness=1)
            box.pack(side='right', fill='x', expand=True, padx=6)
            tk.Label(box, text=title, bg=theme.CARD_BG, fg=theme.SUBTEXT, font=theme.FONT_SMALL).pack(pady=(8, 0))
            lbl = tk.Label(box, textvariable=var, bg=theme.CARD_BG, fg=color, font=(theme.FONT_FAMILY, 18, 'bold'))
            lbl.pack(pady=(0, 10))
            return lbl

        charged_var = tk.StringVar(value='$0.00')
        paid_var = tk.StringVar(value='$0.00')
        balance_var = tk.StringVar(value='$0.00')
        make_card('חויב', charged_var, theme.DARK)
        make_card('שולם', paid_var, theme.SUCCESS)
        balance_lbl = make_card('נשאר לשלם', balance_var, theme.DANGER)

        # --- טופס הוספה ---
        form = ttk.LabelFrame(main, text="רישום חיוב / תשלום", padding=8)
        form.pack(fill='x', pady=6)
        type_var = tk.StringVar(value='charge')
        date_var = tk.StringVar(value=datetime.now().strftime('%Y-%m-%d'))
        amount_var = tk.StringVar()
        note_var = tk.StringVar()

        tk.Label(form, text="סוג:").grid(row=0, column=0, padx=4, pady=4, sticky='e')
        type_cell = tk.Frame(form)
        type_cell.grid(row=0, column=1, padx=4, pady=4, sticky='w')
        tk.Radiobutton(type_cell, text="חיוב (חשבונית)", variable=type_var, value='charge').pack(side='left')
        tk.Radiobutton(type_cell, text="תשלום", variable=type_var, value='payment').pack(side='left', padx=(6, 0))

        tk.Label(form, text="תאריך:").grid(row=0, column=2, padx=4, pady=4, sticky='e')
        date_cell = tk.Frame(form)
        date_cell.grid(row=0, column=3, padx=4, pady=4, sticky='w')
        try:
            from tkcalendar import DateEntry  # type: ignore
            date_entry = DateEntry(date_cell, textvariable=date_var, width=12, date_pattern='yyyy-mm-dd', locale='he_IL')
            try:
                date_entry.set_date(datetime.now())
            except Exception:
                pass
        except Exception:
            date_entry = tk.Entry(date_cell, textvariable=date_var, width=12, justify='center')
        date_entry.pack(side='left')
        if hasattr(self, '_open_date_picker'):
            tk.Button(date_cell, text='📅', width=2, command=lambda: self._open_date_picker(date_entry, date_var, lambda: None)).pack(side='left', padx=(2, 0))

        tk.Label(form, text="סכום $:").grid(row=0, column=4, padx=4, pady=4, sticky='e')
        tk.Entry(form, textvariable=amount_var, width=12, justify='center').grid(row=0, column=5, padx=4, pady=4)
        tk.Label(form, text="הערה:").grid(row=0, column=6, padx=4, pady=4, sticky='e')
        tk.Entry(form, textvariable=note_var, width=34).grid(row=0, column=7, padx=4, pady=4, sticky='we')
        form.columnconfigure(7, weight=1)

        # --- טבלת תנועות ---
        cols = ('date', 'type', 'note', 'net', 'paid', 'balance_net')
        headers = {'date': 'תאריך', 'type': 'סוג', 'note': 'תיאור', 'net': 'חויב', 'paid': 'שולם', 'balance_net': 'יתרה'}
        widths = {'date': 90, 'type': 70, 'note': 380, 'net': 120, 'paid': 120, 'balance_net': 120}
        tree_wrap = tk.Frame(main, bg=theme.PAGE_BG)
        tree_wrap.pack(fill='both', expand=True, pady=(4, 0))
        tree = theme.make_treeview(tree_wrap, columns=cols, show='headings', selectmode='browse')
        for c in cols:
            tree.heading(c, text=headers[c])
            tree.column(c, width=widths[c], anchor='e' if c == 'note' else 'center', stretch=(c == 'note'))
        vs = ttk.Scrollbar(tree_wrap, orient='vertical', command=tree.yview)
        tree.configure(yscrollcommand=vs.set)
        tree.pack(side='left', fill='both', expand=True)
        vs.pack(side='right', fill='y')
        tree.tag_configure('charge', foreground=theme.DARK)
        tree.tag_configure('payment', foreground=theme.SUCCESS)

        def refresh():
            for iid in tree.get_children():
                tree.delete(iid)
            bal = self.data_processor.get_supplier_balance(supplier_id)
            for row in bal['entries']:
                is_charge = row.get('type') == 'charge'
                tree.insert('', 'end', iid=str(row.get('id')), tags=(row.get('type'),), values=(
                    row.get('date', ''),
                    'חיוב' if is_charge else 'תשלום',
                    row.get('note', ''),
                    money(row['net']) if is_charge else '',
                    money(row['paid']) if not is_charge else '',
                    money(row['balance_net']),
                ))
            theme.stripe_tree(tree)
            charged_var.set(money(bal['charged_net']))
            paid_var.set(money(bal['paid']))
            balance_var.set(money(bal['balance_net']))
            balance_lbl.configure(fg=theme.DANGER if bal['balance_net'] > 0.009 else theme.SUCCESS)

        def add_entry():
            try:
                self.data_processor.add_supplier_ledger_entry(
                    supplier_id,
                    type_var.get(),
                    amount_var.get(),
                    date_str=date_var.get().strip(),
                    note=note_var.get(),
                )
            except Exception as e:
                messagebox.showerror("שגיאה", str(e), parent=dialog)
                return
            amount_var.set('')
            note_var.set('')
            refresh()

        def delete_entry():
            sel = tree.selection()
            if not sel:
                messagebox.showinfo("מאזן", "יש לבחור שורה", parent=dialog)
                return
            rec = self.data_processor.get_supplier_ledger_entry(sel[0]) or {}
            label = f"{'חיוב' if rec.get('type') == 'charge' else 'תשלום'} {money(rec.get('amount'))} מ-{rec.get('date', '')}"
            if not messagebox.askyesno("מחיקה", f"למחוק את {label}?", parent=dialog):
                return
            if self.data_processor.delete_supplier_ledger_entry(sel[0]):
                refresh()
            else:
                messagebox.showerror("שגיאה", "לא ניתן למחוק את הרשומה", parent=dialog)

        buttons = tk.Frame(main, bg=theme.PAGE_BG)
        buttons.pack(fill='x', pady=(10, 0))
        theme.make_button(buttons, "הוסף", kind="success", command=add_entry).pack(side='right', padx=4)
        theme.make_button(buttons, "מחק נבחר", kind="danger", command=delete_entry).pack(side='right', padx=4)
        tk.Button(buttons, text="סגור", command=on_close, bg=theme.DARK_2, fg='white', font=(theme.FONT_FAMILY, 10), width=10).pack(side='left', padx=4)
        theme.make_button(
            buttons, "ייצוא PDF", kind="primary",
            command=lambda: self._export_supplier_balance_pdf(supplier, parent=dialog),
        ).pack(side='left', padx=4)
        theme.make_button(
            buttons, "ייצוא Excel", kind="secondary",
            command=lambda: self._export_supplier_balance_excel(supplier, parent=dialog),
        ).pack(side='left', padx=4)

        refresh()

    # ------------------------------------------------------------------
    # ייצוא מאזן ספק (PDF / Excel)
    # ------------------------------------------------------------------
    def _supplier_balance_rows(self, bal: dict) -> list:
        """שורות מאזן ספק בפורמט כרטסת (date/type/id/debit/credit/balance/note)."""
        rows = []
        for e in bal.get('entries') or []:
            is_charge = e.get('type') == 'charge'
            rows.append({
                'date': e.get('date') or '',
                'type': 'חיוב' if is_charge else 'תשלום',
                'id': e.get('id'),
                'source': '',
                'debit': float(e.get('net') or 0) if is_charge else 0,
                'credit': float(e.get('paid') or 0) if not is_charge else 0,
                'balance': float(e.get('balance_net') or 0),
                'note': e.get('note') or '',
            })
        return rows

    def _supplier_export_dir(self) -> str:
        path = os.path.join(os.getcwd(), 'exports', 'supplier_balance')
        os.makedirs(path, exist_ok=True)
        return path

    @staticmethod
    def _supplier_safe_name(name: str) -> str:
        bad = '<>:"/\\|?*'
        clean = ''.join('_' if ch in bad else ch for ch in str(name or '').strip())
        return clean or 'supplier'

    def _export_supplier_balance_pdf(self, supplier: dict, parent=None):
        try:
            from optitex_analyzer.core.baby_basic_note_pdf import generate_baby_basic_ledger_pdf
        except Exception as e:
            messagebox.showerror("שגיאה", f"לא ניתן ליצור PDF:\n{e}", parent=parent)
            return
        supplier_id = int(supplier.get('id'))
        display_name = supplier.get('business_name') or supplier.get('name') or f"ספק {supplier_id}"
        bal = self.data_processor.get_supplier_balance(supplier_id)
        rows = self._supplier_balance_rows(bal)
        if not rows:
            messagebox.showwarning("ייצוא", "אין תנועות במאזן לייצוא.", parent=parent)
            return
        biz_name = ''
        try:
            biz_name = (self.settings.get('business.name', '') or '') if getattr(self, 'settings', None) else ''
        except Exception:
            biz_name = ''
        path = os.path.join(
            self._supplier_export_dir(),
            f"supplier_balance_{self._supplier_safe_name(display_name)}_{datetime.now().strftime('%Y%m%d')}.pdf",
        )
        try:
            generate_baby_basic_ledger_pdf(
                rows,
                path,
                biz_name=biz_name or 'FactorySync',
                partner=display_name,
                source_label=f"סכומים ב-{bal.get('currency') or 'USD'}",
                supplied=bal.get('charged_net'),
                paid=bal.get('paid'),
                balance=bal.get('balance_net'),
                title='מאזן חיובים ותשלומים — ספק',
                partner_label='ספק',
                filter_label='מטבע',
                money_fmt=self._supplier_money,
                show_logo=False,
                show_source=False,
                headers={'debit': 'חויב', 'credit': 'שולם', 'note': 'תיאור'},
                summary_labels={'supplied': 'חויב', 'paid': 'שולם', 'balance': 'נשאר לשלם'},
            )
        except Exception as e:
            messagebox.showerror("שגיאה", f"יצירת PDF נכשלה:\n{e}", parent=parent)
            return
        try:
            os.startfile(path)
        except Exception:
            messagebox.showinfo("PDF", f"הקובץ נשמר:\n{path}", parent=parent)

    def _export_supplier_balance_excel(self, supplier: dict, parent=None):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
            from openpyxl.utils import get_column_letter
        except Exception as e:
            messagebox.showerror("שגיאה", f"נדרש openpyxl לייצוא לאקסל:\n{e}", parent=parent)
            return
        supplier_id = int(supplier.get('id'))
        display_name = supplier.get('business_name') or supplier.get('name') or f"ספק {supplier_id}"
        bal = self.data_processor.get_supplier_balance(supplier_id)
        entries = bal.get('entries') or []
        if not entries:
            messagebox.showwarning("ייצוא", "אין תנועות במאזן לייצוא.", parent=parent)
            return
        currency = bal.get('currency') or 'USD'
        num_fmt = '"$"#,##0.00' if currency.upper() == 'USD' else '#,##0.00'

        wb = Workbook()
        ws = wb.active
        ws.title = 'מאזן'
        ws.sheet_view.rightToLeft = True
        thin = Side(style='thin', color='CCCCCC')
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        head_fill = PatternFill('solid', fgColor='3B82F6')
        head_font = Font(bold=True, color='FFFFFF')

        ws['A1'] = f"מאזן חיובים ותשלומים — {display_name}"
        ws['A1'].font = Font(bold=True, size=14)
        ws['A2'] = f"סכומים ב-{currency}   |   הופק: {datetime.now().strftime('%d/%m/%Y %H:%M')}"
        ws['A2'].font = Font(color='666666')

        headers = ['תאריך', 'סוג', 'מס׳', 'תיאור', 'חויב', 'שולם', 'יתרה']
        hdr_row = 4
        for col, text in enumerate(headers, start=1):
            cell = ws.cell(row=hdr_row, column=col, value=text)
            cell.fill = head_fill
            cell.font = head_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border
        r = hdr_row + 1
        for e in entries:
            is_charge = e.get('type') == 'charge'
            values = [
                e.get('date') or '',
                'חיוב' if is_charge else 'תשלום',
                e.get('id'),
                e.get('note') or '',
                float(e.get('net') or 0) if is_charge else None,
                float(e.get('paid') or 0) if not is_charge else None,
                float(e.get('balance_net') or 0),
            ]
            for col, val in enumerate(values, start=1):
                cell = ws.cell(row=r, column=col, value=val)
                cell.border = border
                if col in (5, 6, 7):
                    cell.number_format = num_fmt
                    cell.alignment = Alignment(horizontal='center')
                elif col == 4:
                    cell.alignment = Alignment(horizontal='right', wrap_text=True, vertical='top')
                else:
                    cell.alignment = Alignment(horizontal='center', vertical='top')
            r += 1

        r += 1
        summary = [
            ('מספר תנועות', len(entries), None),
            ('חויב', float(bal.get('charged_net') or 0), num_fmt),
            ('שולם', float(bal.get('paid') or 0), num_fmt),
            ('נשאר לשלם', float(bal.get('balance_net') or 0), num_fmt),
        ]
        for label, val, fmt in summary:
            ws.cell(row=r, column=1, value=label).font = Font(bold=True)
            c = ws.cell(row=r, column=2, value=val)
            c.font = Font(bold=True)
            if fmt:
                c.number_format = fmt
            r += 1

        widths = {1: 12, 2: 9, 3: 7, 4: 60, 5: 14, 6: 14, 7: 14}
        for col, w in widths.items():
            ws.column_dimensions[get_column_letter(col)].width = w
        ws.freeze_panes = ws.cell(row=hdr_row + 1, column=1)

        path = os.path.join(
            self._supplier_export_dir(),
            f"supplier_balance_{self._supplier_safe_name(display_name)}_{datetime.now().strftime('%Y%m%d')}.xlsx",
        )
        try:
            wb.save(path)
        except Exception as e:
            messagebox.showerror("שגיאה", f"שמירת הקובץ נכשלה (אולי הקובץ פתוח?):\n{e}", parent=parent)
            return
        try:
            os.startfile(path)
        except Exception:
            messagebox.showinfo("Excel", f"הקובץ נשמר:\n{path}", parent=parent)
