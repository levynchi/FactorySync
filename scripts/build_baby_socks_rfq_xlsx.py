# -*- coding: utf-8 -*-
"""RFQ Excel for baby socks 5-pair packs — generic supplier, no price."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
OUT_NAME = "Baby_socks_RFQ_5pair_packs.xlsx"
OUT_XLSX = ROOT / "exports" / OUT_NAME

DATE = date(2026, 9, 14).strftime("%d.%m.%Y")

LINES = [
    {"no": 1, "color": "White", "packs": 1000},
    {"no": 2, "color": "Pink", "packs": 250},
    {"no": 3, "color": "Light Blue", "packs": 250},
]

NAVY = PatternFill("solid", fgColor="1B2A4A")
ALT = PatternFill("solid", fgColor="F3EFE8")
PAPER = PatternFill("solid", fgColor="FAF7F2")
SECTION = PatternFill("solid", fgColor="E8EEF6")
TOTAL = PatternFill("solid", fgColor="1B2A4A")
WHITE = PatternFill("solid", fgColor="FFFFFF")
SWATCH = {
    "White": PatternFill("solid", fgColor="F5F5F4"),
    "Pink": PatternFill("solid", fgColor="F9A8D4"),
    "Light Blue": PatternFill("solid", fgColor="7DD3FC"),
}
THIN = Border(
    left=Side(style="thin", color="D6D3D1"),
    right=Side(style="thin", color="D6D3D1"),
    top=Side(style="thin", color="D6D3D1"),
    bottom=Side(style="thin", color="D6D3D1"),
)
TITLE = Font(name="Calibri", bold=True, size=16, color="1B2A4A")
SUB = Font(name="Calibri", size=11, color="57534E")
HEAD = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
BODY = Font(name="Calibri", size=10, color="1C1917")
BOLD = Font(name="Calibri", bold=True, size=10, color="1C1917")
WHITE_BOLD = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
NOTE = Font(name="Calibri", italic=True, size=9, color="57534E")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
HEADERS = [
    "#",
    "Product",
    "Color",
    "Packs (5 pairs)",
    "Pairs",
    "Unit",
    "Unit price USD / pack",
    "Amount USD",
    "Lead time",
]
NCOLS = len(HEADERS)
QTY_COLS = (4, 5)
PRICE_COLS = (7, 8)


def _paint(cell, *, fill=None, font=None, align=None, border=True) -> None:
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if align is not None:
        cell.alignment = align
    if border:
        cell.border = THIN


def _merge_row(ws: Worksheet, row: int, value: str, font, fill, align=LEFT) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=NCOLS)
    cell = ws.cell(row=row, column=1, value=value)
    _paint(cell, fill=fill, font=font, align=align, border=False)
    for col in range(2, NCOLS + 1):
        _paint(ws.cell(row=row, column=col), fill=fill, border=False)


def build(ws: Worksheet) -> None:
    ws.title = "RFQ"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.oddHeader.left.text = "Aryeh Baby Clothes  ·  Levi Moshe"
    ws.oddHeader.right.text = "RFQ — baby socks"
    ws.oddFooter.left.text = "Confidential — request for quotation"
    ws.oddFooter.right.text = "Page &P of &N"

    widths = [5, 28, 14, 16, 12, 10, 22, 14, 14]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    _merge_row(ws, 1, "REQUEST FOR QUOTATION — BABY SOCKS  (5-PAIR PACKS)", TITLE, PAPER)
    ws.row_dimensions[1].height = 26
    _merge_row(
        ws,
        2,
        f"From Lior  ·  Aryeh Baby Clothes / Levi Moshe    |    To: ______________________________    |    {DATE}",
        SUB,
        PAPER,
    )
    _merge_row(
        ws,
        3,
        "Please quote FOB Turkey or EXW  ·  USD  ·  5 pairs per retail pack  ·  cotton-rich baby socks",
        SUB,
        PAPER,
    )
    _merge_row(
        ws,
        5,
        "Buyer: ARYE MOSHE LEVY  ·  6 Hakishon St, Tel Aviv, Israel  ·  ZIP 6609306  ·  VAT 52064219",
        BODY,
        PAPER,
    )
    _merge_row(
        ws,
        6,
        "Spec: cotton-rich (~90% cotton)  ·  baby sizes 0–12 months  ·  retail 5-pair pack  ·  OEKO-TEX preferred",
        NOTE,
        PAPER,
    )
    _merge_row(
        ws,
        7,
        "Please fill unit price, amount and lead time. Send a sample + packing photo before production. State MOQ and payment terms.",
        NOTE,
        PAPER,
    )

    row = 9
    for col, title in enumerate(HEADERS, 1):
        cell = ws.cell(row=row, column=col, value=title)
        _paint(cell, fill=NAVY, font=HEAD, align=CENTER)
    ws.row_dimensions[row].height = 24
    ws.freeze_panes = "A10"

    total_packs = 0
    total_pairs = 0
    first_data = row + 1
    for i, line in enumerate(LINES):
        row += 1
        packs = line["packs"]
        pairs = packs * 5
        total_packs += packs
        total_pairs += pairs
        fill = ALT if i % 2 else WHITE
        values = [
            line["no"],
            "Baby socks, 5-pair pack",
            line["color"],
            packs,
            pairs,
            "packs",
            None,
            None,
            None,
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            if col == 3:
                _paint(cell, fill=SWATCH[line["color"]], font=BOLD, align=CENTER)
            else:
                _paint(cell, fill=fill, font=BODY, align=CENTER)
            if col in QTY_COLS:
                cell.number_format = "#,##0"
            if col in PRICE_COLS:
                cell.number_format = '"$"#,##0.00'
        ws.row_dimensions[row].height = 22
        amount_cell = ws.cell(row=row, column=8)
        amount_cell.value = f"=D{row}*G{row}"

    last_data = row
    row += 1
    totals = ["", "TOTAL", "", total_packs, total_pairs, "packs", None, None, ""]
    for col, val in enumerate(totals, 1):
        cell = ws.cell(row=row, column=col, value=val)
        _paint(cell, fill=TOTAL, font=WHITE_BOLD, align=CENTER if col != 2 else LEFT)
        if col in QTY_COLS:
            cell.number_format = "#,##0"
        if col in PRICE_COLS:
            cell.number_format = '"$"#,##0.00'
    ws.cell(row=row, column=8).value = f"=SUM(H{first_data}:H{last_data})"
    ws.row_dimensions[row].height = 24

    row += 2
    _merge_row(ws, row, "PLEASE QUOTE", BOLD, SECTION)
    ws.row_dimensions[row].height = 20
    notes = [
        "1. Product: baby socks, packed 5 pairs per retail pack.",
        "2. Colors: White 1,000 packs  ·  Pink 250 packs  ·  Light Blue 250 packs.",
        "3. Total: 1,500 packs  =  7,500 pairs.",
        "4. Composition: cotton-rich, around 90% cotton. Sizes 0–12 months. OEKO-TEX preferred.",
        "5. Please send unit price USD / pack, total amount, lead time, MOQ and payment terms.",
        "6. Please send a sample and a photo of the 5-pair packing before production.",
        "7. Incoterms: FOB Turkey or EXW. Quote in USD.",
    ]
    for text in notes:
        row += 1
        _merge_row(ws, row, text, BODY, PAPER)
        ws.row_dimensions[row].height = 18

    ws.print_area = f"A1:I{row}"


def main() -> None:
    wb = Workbook()
    build(wb.active)
    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_XLSX)
    print(OUT_XLSX)


if __name__ == "__main__":
    main()
