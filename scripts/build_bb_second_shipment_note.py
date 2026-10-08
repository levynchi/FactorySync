# -*- coding: utf-8 -*-
"""Create Baby Basic delivery note for the second Turkey shipment (models 025/204).

Reads the "בייבי בייסיק" sheet of the second-shipment packing list, aggregates pieces
by (model, size), maps to Rivhit barcodes / names and the Arye -> Baby Basic price list
(025 = 8.23 ILS, 204 = 9.24 ILS), saves the note via DataProcessor and writes a PDF.

Run once from the project root:  python scripts/build_bb_second_shipment_note.py
Idempotent: if a note with the same comment already exists it is reused.
"""
from __future__ import annotations

import os
import sys
from collections import OrderedDict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from optitex_analyzer.core.data_processor import DataProcessor  # noqa: E402
from optitex_analyzer.core.baby_basic_note_pdf import generate_baby_basic_note_pdf  # noqa: E402

PACKING = ROOT / "supplier_documents" / "4" / "025-204 second shipment_20261005_080242.xlsx"
SHEET = "בייבי בייסיק"
NOTE_TEXT = "משלוח שני מטורקיה — דגמים 025/204 (26 קרטונים, 1X–18X, 77X–84X). הכל לבן."
NOTE_DATE = datetime.now().strftime("%Y-%m-%d")

# (model, size) -> (barcode, item_name, print_name, unit_price)
ITEMS = {
    (25, "0-3"): ("7297555021786", "בייבי בייסיק בגד גוף שרוול ארוך פלנל לבן (0-3)", "בגד גוף שרוול ארוך", 8.23),
    (25, "3-6"): ("7297555021793", "בייבי בייסיק בגד גוף שרוול ארוך פלנל לבן (3-6)", "בגד גוף שרוול ארוך", 8.23),
    (25, "6-12"): ("7297555021809", "בייבי בייסיק בגד גוף שרוול ארוך פלנל לבן (6-12)", "בגד גוף שרוול ארוך", 8.23),
    (25, "12-18"): ("7297555021816", "בייבי בייסיק בגד גוף שרוול ארוך פלנל לבן (12-18)", "בגד גוף שרוול ארוך", 8.23),
    (25, "18-24"): ("7297555021823", "בייבי בייסיק בגד גוף שרוול ארוך פלנל לבן (18-24)", "בגד גוף שרוול ארוך", 8.23),
    (25, "24-30"): ("7297555021830", "בייבי בייסיק בגד גוף שרוול ארוך פלנל לבן (24-30)", "בגד גוף שרוול ארוך", 8.23),
    (204, "4"): ("7297555022004", "בייבי בייסיק גופיות פלנל (4)", "גופיות פלנל", 9.24),
    (204, "6"): ("7297555022011", "בייבי בייסיק גופיות פלנל (6)", "גופיות פלנל", 9.24),
}
SIZE_ORDER = ["0-3", "3-6", "6-12", "12-18", "18-24", "24-30", "2", "4", "6"]


def load_quantities() -> "OrderedDict[tuple, int]":
    ws = load_workbook(PACKING, data_only=True)[SHEET]
    qty: dict[tuple, int] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        carton, model, size, brand, _packets, pieces = row[:6]
        if model is None or pieces is None or not str(carton or "").strip():
            continue
        try:
            model_i = int(str(model).strip())
            pieces_i = int(float(str(pieces).replace(",", "")))
        except ValueError:
            continue
        if "BABY" not in str(brand or "").upper().replace("İ", "I"):
            continue
        key = (model_i, str(size).strip())
        qty[key] = qty.get(key, 0) + pieces_i
    ordered = OrderedDict(
        sorted(qty.items(), key=lambda kv: (kv[0][0], SIZE_ORDER.index(kv[0][1]) if kv[0][1] in SIZE_ORDER else 99))
    )
    return ordered


def build_lines(qty: "OrderedDict[tuple, int]") -> list[dict]:
    lines = []
    for key, pieces in qty.items():
        if key not in ITEMS:
            raise RuntimeError(f"No barcode/price mapping for model/size {key}")
        barcode, item_name, print_name, price = ITEMS[key]
        lines.append({
            "barcode": barcode,
            "item_name": item_name,
            "print_name": print_name,
            "size": key[1],
            "fabric": "פלנל",
            "color": "לבן",
            "print": "",
            "pack_qty": 5,
            "quantity": pieces,
            "unit_price": price,
        })
    return lines


def main() -> None:
    os.chdir(ROOT)
    dp = DataProcessor()
    qty = load_quantities()
    lines = build_lines(qty)

    existing = next((n for n in dp.get_baby_basic_notes() if n.get("note") == NOTE_TEXT), None)
    if existing:
        note_id = existing["id"]
        print(f"note already exists: #{note_id}")
    else:
        note_id = dp.add_baby_basic_note(
            dp.BABY_BASIC_PARTNER_NAME, NOTE_DATE, lines, note=NOTE_TEXT, source="turkey",
        )
        print(f"created note #{note_id}")

    rec = dp.get_baby_basic_note(note_id)
    for ln in rec["lines"]:
        print(f"  {ln['size']:>6} {ln['barcode']} {ln['quantity']:>5} x {ln['unit_price']:.2f} = {ln['line_total']:,.2f}")
    print(f"total_quantity={rec['total_quantity']} total_amount={rec['total_amount']:,.2f}")
    assert rec["total_quantity"] == 4475, rec["total_quantity"]
    assert abs(rec["total_amount"] - 37874.60) < 0.01, rec["total_amount"]

    out_dir = ROOT / "exports" / "baby_basic_notes"
    out_dir.mkdir(parents=True, exist_ok=True)
    pdf_path = out_dir / f"baby_basic_note_{note_id}_{rec['date']}.pdf"
    generate_baby_basic_note_pdf(rec, str(pdf_path), biz_name="בייבי בייסיק", show_prices=True)
    print(f"pdf={pdf_path}")

    acc = dp.get_baby_basic_account()
    print(f"account: supplied={acc['supplied']:,.2f} paid={acc['paid']:,.2f} balance={acc['balance']:,.2f}")


if __name__ == "__main__":
    main()
