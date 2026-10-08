"""Build an RTL Excel sheet of print quantities per design and color for Hadar."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent.parent / "כמויות_הדפסה_להדר.xlsx"

# name, white (M, L, XL) or None, black (M, L, XL) or None, print note
DESIGNS = [
    ("markio 1", None, (4, 4, 3), "מארקיו מאחורה בכיתוב לבן"),
    ("markio 2", (4, 4, 3), None, "בלי מארקיו מאחורה"),
    ("markio 3", (4, 3, 4), (3, 4, 3), "רקע שחור כיתוב לבן, רקע לבן כיתוב שחור, תווית מאחורה בלי מסגרת"),
    ("markio 4", (4, 3, 4), None, "תווית מרקיו אחורית עם מסגרת"),
    ("markio 5", (4, 3, 4), None, "כיתוב תכלת, מארקיו אחורי שחור"),
    ("markio 6", (3, 4, 3), (4, 3, 4), "מארקיו מקדימה בסגול, מאחורה לא לשים מרקיו"),
    ("markio 7", (3, 4, 4), (3, 4, 3), "בלבן הכיתוב בירוק ובשחור הכיתוב בלבן, בלי דפוס מרקיו מאחורה"),
    ("markio 8", None, (4, 3, 4), "מרקיו מאחורה לבן"),
    ("markio 9", (3, 4, 3), (3, 4, 4), "בלבן הכיתוב בשחור, בלי מרקיו מאחורה"),
    ("markio 10", None, (4, 3, 4), "מרקיו שחור קדמי"),
]

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=12)
TOTAL_FILL = PatternFill("solid", fgColor="D9E1F2")
TOTAL_FONT = Font(bold=True, size=12)
BODY_FONT = Font(size=12)
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center")


def style_sheet(ws, headers, rows, totals, widths, text_cols):
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.append(totals)

    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER
    ws.row_dimensions[1].height = 24

    last = ws.max_row
    for r in range(2, last + 1):
        ws.row_dimensions[r].height = 22
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BORDER
            cell.alignment = RIGHT if c in text_cols else CENTER
            if r == last:
                cell.fill = TOTAL_FILL
                cell.font = TOTAL_FONT
            else:
                cell.font = BODY_FONT
    ws.freeze_panes = "A2"


def build():
    wb = Workbook()

    # Sheet 1: quantities per design and color, no sizes
    ws = wb.active
    ws.title = "כמויות הדפסה"
    rows = []
    for name, white, black, note in DESIGNS:
        w = sum(white) if white else 0
        b = sum(black) if black else 0
        rows.append([name, w or "", b or "", w + b, note])
    total_w = sum(r[1] or 0 for r in rows)
    total_b = sum(r[2] or 0 for r in rows)
    style_sheet(
        ws,
        headers=["עיצוב", "לבן", "שחור", "סה\"כ", "הערות להדפסה"],
        rows=rows,
        totals=["סה\"כ", total_w, total_b, total_w + total_b, ""],
        widths=[14, 10, 10, 10, 70],
        text_cols={1, 5},
    )

    # Sheet 2: breakdown by size
    ws2 = wb.create_sheet("לפי מידות")
    rows2 = []
    for name, white, black, _ in DESIGNS:
        w = white or (0, 0, 0)
        b = black or (0, 0, 0)
        rows2.append([
            name,
            w[0] or "", w[1] or "", w[2] or "", sum(w) or "",
            b[0] or "", b[1] or "", b[2] or "", sum(b) or "",
            sum(w) + sum(b),
        ])
    totals2 = ["סה\"כ"]
    for c in range(1, 10):
        totals2.append(sum((r[c] or 0) for r in rows2))
    style_sheet(
        ws2,
        headers=["עיצוב", "לבן M", "לבן L", "לבן XL", "סה\"כ לבן",
                 "שחור M", "שחור L", "שחור XL", "סה\"כ שחור", "סה\"כ"],
        rows=rows2,
        totals=totals2,
        widths=[14, 9, 9, 9, 11, 9, 9, 9, 11, 10],
        text_cols={1},
    )

    wb.save(OUT)
    print(f"saved {OUT}")


if __name__ == "__main__":
    build()
