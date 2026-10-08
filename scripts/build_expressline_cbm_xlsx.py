# -*- coding: utf-8 -*-
"""Build Express Line CBM / $/CBM workbook for the Import tab."""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
OUT_PATH = ROOT / "import_documents" / "אקספרסליין_חישוב_קוב_2026.xlsx"

NAVY = PatternFill("solid", fgColor="1B2A4A")
SECTION = PatternFill("solid", fgColor="E8EEF6")
ALT = PatternFill("solid", fgColor="F8FAFC")
PAPER = PatternFill("solid", fgColor="FFFFFF")
TOTAL = PatternFill("solid", fgColor="DBEAFE")
GOOD = PatternFill("solid", fgColor="D1FAE5")
WARN = PatternFill("solid", fgColor="FEF3C7")
BAD = PatternFill("solid", fgColor="FEE2E2")
THIN = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)
HEAD = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
BODY = Font(name="Calibri", size=10, color="1C1917")
BOLD = Font(name="Calibri", bold=True, size=10, color="1C1917")
SUB = Font(name="Calibri", size=11, color="475569")
NOTE = Font(name="Calibri", italic=True, size=9, color="64748B")
LEFT = Alignment(horizontal="right", vertical="center", wrap_text=True, readingOrder=2)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
NUM = Alignment(horizontal="center", vertical="center")


def _paint(cell, *, fill=None, font=None, align=None, border=True) -> None:
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if align is not None:
        cell.alignment = align
    if border:
        cell.border = THIN


def _rtl(ws: Worksheet) -> None:
    ws.sheet_view.rightToLeft = True
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4


def _widths(ws: Worksheet, widths: list[float]) -> None:
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _banner(ws: Worksheet, ncols: int, title: str, subtitle: str) -> None:
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(1, 1, title)
    _paint(c, fill=NAVY, font=Font(name="Calibri", bold=True, size=16, color="FFFFFF"), align=LEFT, border=False)
    for col in range(2, ncols + 1):
        _paint(ws.cell(1, col), fill=NAVY, border=False)
    ws.row_dimensions[1].height = 28
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(2, 1, subtitle)
    _paint(c, fill=SECTION, font=SUB, align=LEFT, border=False)
    for col in range(2, ncols + 1):
        _paint(ws.cell(2, col), fill=SECTION, border=False)
    ws.row_dimensions[2].height = 36


def _headers(ws: Worksheet, row: int, titles: list[str]) -> None:
    for col, title in enumerate(titles, 1):
        _paint(ws.cell(row, col, title), fill=NAVY, font=HEAD, align=CENTER)
    ws.row_dimensions[row].height = 28
    ws.auto_filter.ref = f"A{row}:{get_column_letter(len(titles))}{row}"
    ws.freeze_panes = f"A{row + 1}"


def _volume_sheet(ws: Worksheet) -> None:
    ncols = 10
    _rtl(ws)
    _banner(
        ws,
        ncols,
        "חישוב נפח — אקספרסליין 2026",
        "מדידות בפועל: גליל design/פתוח 0.086 קוב; 7 גלילי RIB לבן טובולרי = 70×74×90 ס״מ; קרטון בגדים 60×40×40 ס״מ",
    )
    _widths(ws, [14, 22, 42, 12, 12, 28, 14, 14, 12, 42])

    # Assumptions block
    ws.merge_cells("A4:B4")
    _paint(ws.cell(4, 1, "הנחות מדידה"), fill=SECTION, font=BOLD, align=LEFT)
    _paint(ws.cell(4, 2), fill=SECTION, border=True)
    labels = [
        (5, "קוב לגליל פתוח / design", 0.086, "קוב", "25 גלילי אפריל + 17 גלילי design במאי"),
        (6, "אורך חבילת 7 גלילי RIB לבן (מ׳)", 0.70, "מ׳", "מדידה ל-7 גלילים ביחד, לא לגליל בודד"),
        (7, "רוחב חבילת 7 גלילי RIB לבן (מ׳)", 0.74, "מ׳", ""),
        (8, "גובה חבילת 7 גלילי RIB לבן (מ׳)", 0.90, "מ׳", ""),
        (9, "מספר גלילים במדידה", 7, "גלילים", ""),
        (10, "אורך קרטון בגדים (מ׳)", 0.60, "מ׳", "משלוח עומר 9.9.2026"),
        (11, "רוחב קרטון בגדים (מ׳)", 0.40, "מ׳", ""),
        (12, "גובה קרטון בגדים (מ׳)", 0.40, "מ׳", ""),
        (13, "קוב ייחוס לבדים", 280, "$/קוב", "ממוצע אפריל 268 + מאי 289"),
    ]
    for row, label, value, unit, note in labels:
        _paint(ws.cell(row, 1, label), font=BODY, align=LEFT)
        cell = ws.cell(row, 2, value)
        if isinstance(value, float) and value < 2:
            cell.number_format = "0.000"
        elif isinstance(value, float):
            cell.number_format = "0.00"
        _paint(cell, font=BOLD, align=NUM, fill=WARN)
        _paint(ws.cell(row, 3, unit), font=BODY, align=CENTER)
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=10)
        _paint(ws.cell(row, 4, note), font=NOTE, align=LEFT, border=False)
        for col in range(5, 11):
            _paint(ws.cell(row, col), border=False)

    ws.cell(14, 2, "=B6*B7*B8")
    ws.cell(14, 2).number_format = "0.0000"
    _paint(ws.cell(14, 1, "נפח 7 גלילי RIB לבן"), font=BOLD, align=LEFT, fill=TOTAL)
    _paint(ws.cell(14, 2), font=BOLD, align=NUM, fill=TOTAL)
    _paint(ws.cell(14, 3, "קוב"), font=BOLD, align=CENTER, fill=TOTAL)
    ws.merge_cells("D14:J14")
    _paint(ws.cell(14, 4, "0.70 × 0.74 × 0.90"), font=NOTE, align=LEFT, fill=TOTAL)
    for col in range(5, 11):
        _paint(ws.cell(14, col), fill=TOTAL, border=False)

    ws.cell(15, 2, "=B14/B9")
    ws.cell(15, 2).number_format = "0.0000"
    _paint(ws.cell(15, 1, "קוב לגליל RIB לבן טובולרי"), font=BOLD, align=LEFT, fill=GOOD)
    _paint(ws.cell(15, 2), font=BOLD, align=NUM, fill=GOOD)
    _paint(ws.cell(15, 3, "קוב"), font=BOLD, align=CENTER, fill=GOOD)
    ws.merge_cells("D15:J15")
    _paint(ws.cell(15, 4, "≈ 0.0666 קוב לגליל"), font=NOTE, align=LEFT, fill=GOOD)
    for col in range(5, 11):
        _paint(ws.cell(15, col), fill=GOOD, border=False)

    headers = [
        "תאריך",
        "משלוח",
        "פירוט",
        "כמות",
        "יחידה",
        "מידות / בסיס",
        "קוב ליחידה",
        "קוב סה״כ",
        "ק״ג",
        "הערות",
    ]
    _headers(ws, 17, headers)

    # Row 18 April
    r = 18
    _paint(ws.cell(r, 1, "2026-04-09"), font=BODY, align=CENTER)
    _paint(ws.cell(r, 2, "בדים — גלילים פתוחים"), font=BODY, align=LEFT)
    _paint(ws.cell(r, 3, "כל 25 הגלילים עם design / פתוחים"), font=BODY, align=LEFT)
    _paint(ws.cell(r, 4, 25), font=BODY, align=NUM)
    _paint(ws.cell(r, 5, "גליל"), font=BODY, align=CENTER)
    _paint(ws.cell(r, 6, "0.086 קוב לגליל (מדידה)"), font=BODY, align=LEFT)
    c = ws.cell(r, 7, "=B5")
    c.number_format = "0.000"
    _paint(c, font=BODY, align=NUM)
    c = ws.cell(r, 8, "=D18*G18")
    c.number_format = "0.000"
    _paint(c, font=BOLD, align=NUM)
    _paint(ws.cell(r, 9, 502), font=BODY, align=NUM)
    _paint(ws.cell(r, 10, "packing list 30.03.2026; בדרישה נרשם 9 קוב (=9 יחידות)"), font=NOTE, align=LEFT)

    # Row 19 May design
    r = 19
    _paint(ws.cell(r, 1, "2026-05-10"), font=BODY, align=CENTER, fill=ALT)
    _paint(ws.cell(r, 2, "בדים — design"), font=BODY, align=LEFT, fill=ALT)
    _paint(ws.cell(r, 3, "הזמנות 97899 + 97898 (R-A1888)"), font=BODY, align=LEFT, fill=ALT)
    _paint(ws.cell(r, 4, 17), font=BODY, align=NUM, fill=ALT)
    _paint(ws.cell(r, 5, "גליל"), font=BODY, align=CENTER, fill=ALT)
    _paint(ws.cell(r, 6, "אותו נפח כמו גליל פתוח"), font=BODY, align=LEFT, fill=ALT)
    c = ws.cell(r, 7, "=B5")
    c.number_format = "0.000"
    _paint(c, font=BODY, align=NUM, fill=ALT)
    c = ws.cell(r, 8, "=D19*G19")
    c.number_format = "0.000"
    _paint(c, font=BOLD, align=NUM, fill=ALT)
    _paint(ws.cell(r, 9, None), font=BODY, align=NUM, fill=ALT)
    _paint(ws.cell(r, 10, "17 גלילים עם design מתוך 43"), font=NOTE, align=LEFT, fill=ALT)

    # Row 20 May white tubular
    r = 20
    _paint(ws.cell(r, 1, "2026-05-10"), font=BODY, align=CENTER)
    _paint(ws.cell(r, 2, "בדים — RIB לבן טובולרי"), font=BODY, align=LEFT)
    _paint(ws.cell(r, 3, "הזמנה 98697; מדידה ל-7 גלילים מייצגים"), font=BODY, align=LEFT)
    _paint(ws.cell(r, 4, 26), font=BODY, align=NUM)
    _paint(ws.cell(r, 5, "גליל"), font=BODY, align=CENTER)
    _paint(ws.cell(r, 6, "70×74×90 ס״מ ל-7 גלילים"), font=BODY, align=LEFT)
    c = ws.cell(r, 7, "=B15")
    c.number_format = "0.0000"
    _paint(c, font=BODY, align=NUM)
    c = ws.cell(r, 8, "=D20*G20")
    c.number_format = "0.000"
    _paint(c, font=BOLD, align=NUM)
    _paint(ws.cell(r, 9, None), font=BODY, align=NUM)
    _paint(ws.cell(r, 10, "לא לגליל בודד — הנפח הוא ל-7 גלילים ביחד"), font=NOTE, align=LEFT)

    # Row 21 May total
    r = 21
    _paint(ws.cell(r, 1, "2026-05-10"), font=BOLD, align=CENTER, fill=TOTAL)
    _paint(ws.cell(r, 2, "סה״כ משלוח מאי"), font=BOLD, align=LEFT, fill=TOTAL)
    _paint(ws.cell(r, 3, "17 design + 26 לבן טובולרי"), font=BOLD, align=LEFT, fill=TOTAL)
    c = ws.cell(r, 4, "=D19+D20")
    _paint(c, font=BOLD, align=NUM, fill=TOTAL)
    _paint(ws.cell(r, 5, "גליל"), font=BOLD, align=CENTER, fill=TOTAL)
    _paint(ws.cell(r, 6, ""), font=BODY, align=LEFT, fill=TOTAL)
    _paint(ws.cell(r, 7, ""), font=BODY, align=NUM, fill=TOTAL)
    c = ws.cell(r, 8, "=H19+H20")
    c.number_format = "0.00"
    _paint(c, font=BOLD, align=NUM, fill=TOTAL)
    _paint(ws.cell(r, 9, 804.1), font=BOLD, align=NUM, fill=TOTAL)
    ws.cell(r, 9).number_format = "0.0"
    _paint(ws.cell(r, 10, "packing list 14.5.2026 / דרישה #1581"), font=NOTE, align=LEFT, fill=TOTAL)

    # Row 22 September
    r = 22
    _paint(ws.cell(r, 1, "2026-09-09"), font=BODY, align=CENTER, fill=ALT)
    _paint(ws.cell(r, 2, "בגדים — קרטונים"), font=BODY, align=LEFT, fill=ALT)
    _paint(ws.cell(r, 3, "משלוח עומר 200/201/202"), font=BODY, align=LEFT, fill=ALT)
    _paint(ws.cell(r, 4, 53), font=BODY, align=NUM, fill=ALT)
    _paint(ws.cell(r, 5, "קרטון"), font=BODY, align=CENTER, fill=ALT)
    _paint(ws.cell(r, 6, "60 × 40 × 40 ס״מ"), font=BODY, align=LEFT, fill=ALT)
    c = ws.cell(r, 7, "=B10*B11*B12")
    c.number_format = "0.000"
    _paint(c, font=BODY, align=NUM, fill=ALT)
    c = ws.cell(r, 8, "=D22*G22")
    c.number_format = "0.00"
    _paint(c, font=BOLD, align=NUM, fill=ALT)
    _paint(ws.cell(r, 9, None), font=BODY, align=NUM, fill=ALT)
    _paint(ws.cell(r, 10, "דרישה #2432 — 53 קרטונים × 42$"), font=NOTE, align=LEFT, fill=ALT)

    ws.merge_cells("A24:J26")
    summary = (
        "סיכום נפח בפועל: אפריל 2.15 קוב (25×0.086) · מאי 3.19 קוב (1.462+1.732) · ספטמבר 5.09 קוב (53×0.096). "
        "בדרישת אפריל נרשם 9 קוב — זה מספר היחידות שחויבו (9×64$), לא הנפח האמיתי."
    )
    _paint(ws.cell(24, 1, summary), fill=SECTION, font=BODY, align=LEFT, border=False)
    for col in range(2, 11):
        _paint(ws.cell(24, col), fill=SECTION, border=False)
        _paint(ws.cell(25, col), fill=SECTION, border=False)
        _paint(ws.cell(26, col), fill=SECTION, border=False)
    ws.row_dimensions[24].height = 22
    ws.row_dimensions[25].height = 22
    ws.row_dimensions[26].height = 22

    ws.print_title_rows = "1:2"
    ws.page_setup.fitToHeight = 1
    ws.oddFooter.right.text = "עמוד &P מתוך &N"
    ws.oddFooter.left.text = "אריה בגדי תינוקות — חישוב קוב אקספרסליין"


def _price_sheet(ws: Worksheet) -> None:
    ncols = 16
    _rtl(ws)
    _banner(
        ws,
        ncols,
        "מחיר לקוב — אקספרסליין 2026",
        "הובלה בלבד (ללא שווי סחורה). מאי: בדרישה #1581 נרשם 1,720$ (43×40) — בפועל חויב 922$ (≈21.5$×43). ייחוס בדים ≈ 280 $/קוב.",
    )
    _widths(ws, [13, 22, 14, 12, 13, 12, 16, 16, 14, 10, 12, 14, 14, 12, 14, 36])

    headers = [
        "תאריך",
        "משלוח",
        "יחידות בדרישה",
        "$ ליחידה",
        "הובלה $",
        "קוב בפועל",
        "$ לקוב לפני מע״מ",
        "$ לקוב כולל 18%",
        "₪ לקוב לפני מע״מ",
        "ק״ג",
        "$ לק״ג",
        "ייחוס $/קוב",
        "הובלה צפויה $",
        "פער $",
        "שער ₪/$",
        "הערות",
    ]
    _headers(ws, 4, headers)

    # April — volume from volume sheet H18, freight 9*64=576, rate 3.14, kg 502
    r = 5
    fill = PAPER
    _paint(ws.cell(r, 1, "2026-04-09"), font=BODY, align=CENTER, fill=fill)
    _paint(ws.cell(r, 2, "בדים — גלילים פתוחים"), font=BODY, align=LEFT, fill=fill)
    _paint(ws.cell(r, 3, 9), font=BODY, align=NUM, fill=fill)
    _paint(ws.cell(r, 4, 64), font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 5, "=C5*D5")
    c.number_format = "#,##0.00"
    _paint(c, font=BOLD, align=NUM, fill=fill)
    c = ws.cell(r, 6, "='חישוב נפח'!H18")
    c.number_format = "0.00"
    _paint(c, font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 7, "=E5/F5")
    c.number_format = "0.00"
    _paint(c, font=BOLD, align=NUM, fill=fill)
    c = ws.cell(r, 8, "=G5*1.18")
    c.number_format = "0.00"
    _paint(c, font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 9, "=E5*O5/F5")
    c.number_format = "#,##0.00"
    _paint(c, font=BODY, align=NUM, fill=fill)
    _paint(ws.cell(r, 10, 502), font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 11, "=E5/J5")
    c.number_format = "0.00"
    _paint(c, font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 12, "='חישוב נפח'!B13")
    c.number_format = "0"
    _paint(c, font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 13, "=F5*L5")
    c.number_format = "#,##0.00"
    _paint(c, font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 14, "=E5-M5")
    c.number_format = "#,##0.00"
    _paint(c, font=BODY, align=NUM, fill=fill)
    _paint(ws.cell(r, 15, 3.14), font=BODY, align=NUM, fill=fill)
    ws.cell(r, 15).number_format = "0.00"
    _paint(ws.cell(r, 16, "9×64$ בדרישה; נפח אמיתי 2.15 קוב → 268 $/קוב, 1.15 $/ק״ג"), font=NOTE, align=LEFT, fill=fill)

    # May — actual 922 not 1720
    r = 6
    fill = ALT
    _paint(ws.cell(r, 1, "2026-05-10"), font=BODY, align=CENTER, fill=fill)
    _paint(ws.cell(r, 2, "בדים — RIB (דרישה #1581)"), font=BODY, align=LEFT, fill=fill)
    _paint(ws.cell(r, 3, 43), font=BODY, align=NUM, fill=fill)
    _paint(ws.cell(r, 4, 21.5), font=BODY, align=NUM, fill=fill)
    ws.cell(r, 4).number_format = "0.0"
    c = ws.cell(r, 5, 922)
    c.number_format = "#,##0.00"
    _paint(c, font=BOLD, align=NUM, fill=fill)
    c = ws.cell(r, 6, "='חישוב נפח'!H21")
    c.number_format = "0.00"
    _paint(c, font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 7, "=E6/F6")
    c.number_format = "0.00"
    _paint(c, font=BOLD, align=NUM, fill=fill)
    c = ws.cell(r, 8, "=G6*1.18")
    c.number_format = "0.00"
    _paint(c, font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 9, "=E6*O6/F6")
    c.number_format = "#,##0.00"
    _paint(c, font=BODY, align=NUM, fill=fill)
    _paint(ws.cell(r, 10, 804.1), font=BODY, align=NUM, fill=fill)
    ws.cell(r, 10).number_format = "0.0"
    c = ws.cell(r, 11, "=E6/J6")
    c.number_format = "0.00"
    _paint(c, font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 12, "='חישוב נפח'!B13")
    c.number_format = "0"
    _paint(c, font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 13, "=F6*L6")
    c.number_format = "#,##0.00"
    _paint(c, font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 14, "=E6-M6")
    c.number_format = "#,##0.00"
    _paint(c, font=BODY, align=NUM, fill=fill)
    _paint(ws.cell(r, 15, 3.00), font=BODY, align=NUM, fill=fill)
    ws.cell(r, 15).number_format = "0.00"
    _paint(
        ws.cell(r, 16, "בפועל 922$ (לא 1,720$ שבדרישה). 289 $/קוב, 1.15 $/ק״ג. 43×21.5=924.5 ≈ 922"),
        font=NOTE,
        align=LEFT,
        fill=fill,
    )

    # September
    r = 7
    fill = PAPER
    _paint(ws.cell(r, 1, "2026-09-09"), font=BODY, align=CENTER, fill=fill)
    _paint(ws.cell(r, 2, "בגדים — קרטונים (#2432)"), font=BODY, align=LEFT, fill=fill)
    _paint(ws.cell(r, 3, 53), font=BODY, align=NUM, fill=fill)
    _paint(ws.cell(r, 4, 42), font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 5, "=C7*D7")
    c.number_format = "#,##0.00"
    _paint(c, font=BOLD, align=NUM, fill=fill)
    c = ws.cell(r, 6, "='חישוב נפח'!H22")
    c.number_format = "0.00"
    _paint(c, font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 7, "=E7/F7")
    c.number_format = "0.00"
    _paint(c, font=BOLD, align=NUM, fill=fill)
    c = ws.cell(r, 8, "=G7*1.18")
    c.number_format = "0.00"
    _paint(c, font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 9, "=E7*O7/F7")
    c.number_format = "#,##0.00"
    _paint(c, font=BODY, align=NUM, fill=fill)
    _paint(ws.cell(r, 10, None), font=BODY, align=NUM, fill=fill)
    _paint(ws.cell(r, 11, None), font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 12, "='חישוב נפח'!B13")
    c.number_format = "0"
    _paint(c, font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 13, "=F7*L7")
    c.number_format = "#,##0.00"
    _paint(c, font=BODY, align=NUM, fill=fill)
    c = ws.cell(r, 14, "=E7-M7")
    c.number_format = "#,##0.00"
    _paint(c, font=BOLD, align=NUM, fill=BAD)
    _paint(ws.cell(r, 15, 3.02), font=BODY, align=NUM, fill=fill)
    ws.cell(r, 15).number_format = "0.00"
    _paint(
        ws.cell(r, 16, "42$ לקרטון = 437 $/קוב. לפי 280 $/קוב צפוי ~27$ לקרטון / ~1,425$. פער ~800$"),
        font=NOTE,
        align=LEFT,
        fill=fill,
    )

    # Totals row
    r = 8
    _paint(ws.cell(r, 1, ""), font=BOLD, align=CENTER, fill=TOTAL)
    _paint(ws.cell(r, 2, "סה״כ הובלה $"), font=BOLD, align=LEFT, fill=TOTAL)
    for col in range(3, 5):
        _paint(ws.cell(r, col), fill=TOTAL)
    c = ws.cell(r, 5, "=E5+E6+E7")
    c.number_format = "#,##0.00"
    _paint(c, font=BOLD, align=NUM, fill=TOTAL)
    c = ws.cell(r, 6, "=F5+F6+F7")
    c.number_format = "0.00"
    _paint(c, font=BOLD, align=NUM, fill=TOTAL)
    c = ws.cell(r, 7, "=E8/F8")
    c.number_format = "0.00"
    _paint(c, font=BOLD, align=NUM, fill=TOTAL)
    for col in range(8, 13):
        _paint(ws.cell(r, col), fill=TOTAL)
    c = ws.cell(r, 13, "=M5+M6+M7")
    c.number_format = "#,##0.00"
    _paint(c, font=BOLD, align=NUM, fill=TOTAL)
    c = ws.cell(r, 14, "=N5+N6+N7")
    c.number_format = "#,##0.00"
    _paint(c, font=BOLD, align=NUM, fill=TOTAL)
    _paint(ws.cell(r, 15), fill=TOTAL)
    _paint(ws.cell(r, 16, "ממוצע משוקלל $/קוב על שלושת המשלוחים"), font=NOTE, align=LEFT, fill=TOTAL)

    # Comparison box
    ws.merge_cells("A10:P10")
    _paint(ws.cell(10, 1, "מה המספרים אומרים"), fill=SECTION, font=BOLD, align=LEFT, border=False)
    for col in range(2, 17):
        _paint(ws.cell(10, col), fill=SECTION, border=False)

    lines = [
        "משלוחי הבדים (אפריל+מאי) יצאו כמעט אותו מחיר לק״ג (1.15 $/ק״ג) ולקוב (~268–289 $/קוב). סביר שאקספרסליין מתמחר בדים לפי משקל/נפח, לא לפי קרטון.",
        "במאי הדרישה #1581 רשמה 40$ לקרטון (1,720$) — זה היה טעות. החיוב בפועל 922$ מיישר את המחיר לקוב עם אפריל.",
        "משלוח הבגדים בספטמבר חויב 42$ לקרטון = 437 $/קוב — פי ~1.6 ממחיר הבדים לקוב. לפי 280 $/קוב קרטון 60×40×40 היה אמור לעלות ~27$, לא 42$.",
        "פער ספטמבר: 2,226$ ששולם מול ~1,425$ צפוי = כ-800$ (כ-2,400 ₪ בשער 3.02). שווה לברר מולם אם הבגדים מתומחרים אחרת מהבד.",
        "דרישת #1581 לאחר תיקון הובלה: בדים 5,533$ + הובלה 922$ = 6,455$ + מע״מ 18% = 7,616.90$ ≈ 22,850.70 ₪ (שער 3.00). המסמך המקורי עדיין מציג 8,558.54$.",
    ]
    for i, line in enumerate(lines):
        row = 11 + i
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        _paint(ws.cell(row, 1, line), font=BODY, align=LEFT, border=False)
        ws.row_dimensions[row].height = 22
        for col in range(2, ncols + 1):
            _paint(ws.cell(row, col), border=False)

    ws.print_title_rows = "1:4"
    ws.page_setup.fitToHeight = 1
    ws.oddFooter.right.text = "עמוד &P מתוך &N"
    ws.oddFooter.left.text = "אריה בגדי תינוקות — מחיר לקוב אקספרסליין"


def build() -> Path:
    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_vol = wb.active
    ws_vol.title = "חישוב נפח"
    ws_price = wb.create_sheet("מחיר לקוב")
    _volume_sheet(ws_vol)
    _price_sheet(ws_price)
    wb.save(OUT_PATH)
    print(f"wrote {OUT_PATH}")
    return OUT_PATH


if __name__ == "__main__":
    build()
