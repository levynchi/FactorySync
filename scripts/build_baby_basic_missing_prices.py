# -*- coding: utf-8 -*-
"""
השלמת מחירון בייבי בייסיק 2026-27 — מוצרים לבנים שאינם במחירון
"מחירון מעודכן אריה לבייבי בייסיק 2026 2027 .xlsx".

שיטת חישוב (זהה לכלל "מכירה לבייבי בייסיק = עלות + 20%" מטאב טורקיה במחירון):
    עלות = בד (שטח גזרה מ"ר × מחיר למ"ר) + טיקטקים (0.30 ליח') + גומי (0.50) + סרט (0.50) + תפירה (לפי דגם)
    מחיר לבייבי בייסיק = עלות × 1.20, עיגול ל-0.1 ₪
מקורות:
    - "מחירי עלות אריה.xlsx" (מייל 10.08.2026) — עלויות תפירה/טיקטקים/בד לפלנל לבן.
    - products_catalog.json — שטחי גזרה (מ"ר) לכל דגם/מידה; fabric_prices.json — פלנל 11 ₪/מ"ר, טריקו 10 ₪/מ"ר.
    - rivhit_products.json — מחיר חנויות של אריה (לייחוס בלבד).
פריטים בלי נתוני גזרה מסומנים "הערכה" / "חסר נתוני עלות".

הרצה:  python scripts/build_baby_basic_missing_prices.py
פלט:   exports/baby_basic_prices/מחירון_השלמה_לבן_2026_2027.xlsx  (+ צירוף להתכתבות המחירון בטאב בייבי בייסיק)
"""
from __future__ import annotations

import json
import os
import sys
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
os.chdir(ROOT)
sys.path.insert(0, ROOT)

OUT_DIR = os.path.join('exports', 'baby_basic_prices')
OUT_FILE = os.path.join(OUT_DIR, 'מחירון_השלמה_לבן_2026_2027.xlsx')
PRICE_LIST_THREAD_ID = '19ecf70f4393432a'   # Gmail thread: "מחירון חורף 2026-2027" (10.08.2026)

MARGIN = 1.20
FABRIC_PRICE = {'פלנל': 11.0, 'טריקו': 10.0}   # ₪ למ"ר (fabric_prices.json)
TICK, ELASTIC, RIBBON = 0.30, 0.50, 0.50          # item_cost_settings / מחירי עלות אריה
SEWING = {                                        # ₪ ליחידה, מתוך "מחירי עלות אריה.xlsx"
    'בגד גוף ש.א': 2.75, 'בגד גוף מעטפת': 3.10, 'רגליות': 1.75, 'גטקס': 1.95,
    'גופיות פלנל': 3.10, 'חולצת בייבי': 2.75, 'אוברול': 4.10,
}

# מחירים קיימים במחירון 2026-27 (לייחוס/אקסטרפולציה)
EXCEL_OVERALL = {'0-3': 12.5, '3-6': 12.7, '6-12': 13.5}
EXCEL_SABA = {'3': 5.5, '4': 5.7, '5': 5.9, '6': 6.0, '8': 6.2, '10': 6.4}

# שטחי גזרה (מ"ר) — products_catalog.json
SQM = {
    'בגד גוף ש.א': {'0-3': 0.231, '3-6': 0.264, '6-12': 0.299, '12-18': 0.337, '18-24': 0.376, '24-30': 0.417},
    'בגד גוף מעטפת': {'0-3': 0.264},
    'רגליות': {'0-3': 0.195, '3-6': 0.218, '6-12': 0.238, '12-18': 0.273, '18-24': 0.310, '24-30': 0.349},
    'גטקס': {'3-6': 0.219, '6-12': 0.252, '12-18': 0.287, '18-24': 0.325, '24-30': 0.364},
    'אוברול': {'0-3': 0.407, '3-6': 0.466, '6-12': 0.528, '12-18': 0.584},
    'גופיות סבא': {'3': 0.176, '4': 0.203, '5': 0.232, '6': 0.263, '8': 0.296, '10': 0.327,
                   '12': 0.358, '14': 0.391, '16': 0.424},
}
TICKS = {'בגד גוף ש.א': {'0-3': 3, '3-6': 3, '6-12': 3, '12-18': 4, '18-24': 4, '24-30': 4},
         'בגד גוף מעטפת': {'0-3': 5}, 'אוברול': {'*': 1}}

# עלות כוללת ישירות מ"מחירי עלות אריה.xlsx" (פלנל לבן, ייצור בארץ)
COST_SHEET = {
    ('גופיות פלנל', '12'): 11.67, ('גופיות פלנל', '14'): 12.65, ('גופיות פלנל', '16'): 13.68,
    ('גופיות פלנל', 'S'): 15.48, ('גופיות פלנל', 'M'): 15.99, ('גופיות פלנל', 'L'): 16.52,
    ('גופיות פלנל', 'XL'): 17.32, ('גופיות פלנל', 'XXL'): 18.14,
    ('חולצת בייבי', '3-6'): 6.67, ('חולצת בייבי', '6-12'): 7.02, ('חולצת בייבי', '12-18'): 7.39,
    ('חולצת בייבי', '18-24'): 7.77, ('חולצת בייבי', '24-30'): 8.17,
}

# מחירי חנויות בריווחית (item_sale_nis) — לייחוס
RIVHIT_SALE = {
    'גופיות פלנל 12': 15.2, 'גופיות פלנל 14': 16.4, 'גופיות פלנל 16': 16.4, 'גופיות פלנל S': 21.9,
    'חולצת בייבי': None, 'אוברול פלנל': 26.0, 'אוברול טריקו': 26.0,
    'בגד גוף טריקו ארוך': 12.8, 'בגד גוף מעטפת טריקו': 13.0, 'רגליות טריקו': 9.5, 'רגליות טריקו 3-6': 9.42,
    'גטקס טריקו': 8.6, 'גופיות סבא 2': 8.0, 'גופיות סבא 12': 10.5, 'גופיות סבא 14': 11.5, 'גופיות סבא 16': 11.5,
    'גופיות בייב קצר': 9.0, 'חזיות סגור פלנל': 10.1, 'חזיות טיקטק פלנל': 11.3, 'חזיות סגור טריקו': 8.0,
    'חזיות טיקטק טריקו': 9.5, 'סט חזיה+רגלית טריקו': 18.5, 'רגליות פלנל N.B': 10.7, 'תחתונים': 6.0,
    'חיתולי פלנל': 25.0, 'סדין+ציפה פלנל': 63.8, 'גטקס פלנל': 11.4,
}
# יחס ממוצע מחיר בייבי בייסיק / מחיר חנויות על הפריטים שכבר מתומחרים (בגד גוף 0.64, רגליות 0.56,
# גטקס 0.64, מעטפת 0.70, גופיות בייבי 0.75, גופיות פלנל 0.72, גופיות סבא 0.65)
STORE_RATIO = 0.66


def r1(x: float) -> float:
    return round(x + 1e-9, 1)


def calc(model: str, size: str, fabric: str, ticks: int = 0, elastic: int = 0, ribbon: int = 0,
         sewing_model: str | None = None, sqm: float | None = None):
    """עלות מחושבת לפי בד+נלווים+תפירה. מחזיר (עלות, פירוט)."""
    area = sqm if sqm is not None else SQM[model][size]
    fab = area * FABRIC_PRICE[fabric]
    sew = SEWING[sewing_model or model]
    acc = ticks * TICK + elastic * ELASTIC + ribbon * RIBBON
    cost = fab + acc + sew
    detail = f"בד {area:.3f} מ\"ר×{FABRIC_PRICE[fabric]:.0f}={fab:.2f} + נלווים {acc:.2f} + תפירה {sew:.2f}"
    return cost, detail


rows: list[dict] = []


def add(cat, name, size, fabric, cost, price, basis, store=None, note='', status='מחושב'):
    rows.append(dict(cat=cat, name=name, size=size, fabric=fabric, cost=cost, price=price,
                     basis=basis, store=store, note=note, status=status))


# ---------------------------------------------------------------- חורף — פלנל לבן
CAT_W = 'חורף — פלנל לבן'
for size in ['12', '14', '16', 'S', 'M', 'L', 'XL', 'XXL']:
    c = COST_SHEET[('גופיות פלנל', size)]
    add(CAT_W, 'גופיות פלנל', size, 'פלנל', c, r1(c * MARGIN), 'מחירי עלות אריה.xlsx',
        RIVHIT_SALE.get(f'גופיות פלנל {size}'),
        'באפליקציה מידה 12 רשומה כיום 11.8 (הועתק ממידה 10)' if size == '12' else
        ('מידות מבוגרים' if size in ('S', 'M', 'L', 'XL', 'XXL') else ''))

for size in ['3-6', '6-12', '12-18', '18-24', '24-30']:
    c = COST_SHEET[('חולצת בייבי', size)]
    add(CAT_W, 'חולצת בייבי פלנל (טיקטקים בכתף)', size, 'פלנל', c, r1(c * MARGIN), 'מחירי עלות אריה.xlsx', None)

# אוברול פלנל 12-18 — אקסטרפולציה מהמחירון (0-3/3-6/6-12) לפי יחס עלות הבד+תפירה
c612, _ = calc('אוברול', '6-12', 'פלנל', ticks=1)
c1218, d1218 = calc('אוברול', '12-18', 'פלנל', ticks=1)
p1218 = r1(EXCEL_OVERALL['6-12'] * (c1218 / c612))
add(CAT_W, 'אוברול פלנל (רוכסן)', '12-18', 'פלנל', c1218, p1218,
    f'המשך סדרת המחירון (6-12 = 13.5) לפי יחס עלות; {d1218}', RIVHIT_SALE['אוברול פלנל'],
    'במחירון יש רק 0-3 / 3-6 / 6-12', 'הערכה')

# גטקס פלנל ילדים 4 / 6 — אין שטח גזרה בקטלוג; הערכה לפי המשך הסדרה (24-30 = 0.364 מ"ר)
for size, area in (('4', 0.46), ('6', 0.56)):
    c, d = calc('גטקס', size, 'פלנל', elastic=1, sqm=area)
    add(CAT_W, 'גטקס פתוח פלנל ילדים', size, 'פלנל', c, r1(c * MARGIN), f'שטח גזרה משוער {area} מ"ר; {d}',
        RIVHIT_SALE['גטקס פלנל'], 'יש למדוד גזרה באופטיטקס לאימות', 'הערכה')

# פגים / N.B פלנל
c, d = calc('רגליות', 'פגים', 'פלנל', elastic=1, sqm=0.17)
add(CAT_W, 'רגליות פלנל פגים (N.B)', 'פגים', 'פלנל', c, r1(c * MARGIN), f'שטח משוער 0.17 מ"ר; {d}',
    RIVHIT_SALE['רגליות פלנל N.B'], 'ברקוד 7297555020871', 'הערכה')
c, d = calc('בגד גוף מעטפת', 'פגים', 'פלנל', ticks=5, ribbon=1, sqm=0.23)
add(CAT_W, 'בגד גוף מעטפת פלנל פגים (N.B)', 'פגים', 'פלנל', c, r1(c * MARGIN), f'שטח משוער 0.23 מ"ר; {d}',
    None, 'ברקוד 7297555020888', 'הערכה')
for name, key in (('חזיות סגור פלנל N.B', 'חזיות סגור פלנל'), ('חזיות טיקטק פלנל + כפפה N.B', 'חזיות טיקטק פלנל')):
    s = RIVHIT_SALE[key]
    add(CAT_W, name, 'N.B', 'פלנל', None, r1(s * STORE_RATIO), f'מחיר חנויות {s} × {STORE_RATIO}', s,
        'אין גזרה בקטלוג — תומחר לפי יחס ממוצע מחירון/חנויות', 'הערכה')

for size in ['0-3', '3-6', '6-12', '12-18']:
    add(CAT_W, 'אוברול ארוך סואגי', size, 'פלנל', None, None, '—', None,
        'בקטלוג רק כמות טיקטקים (9-10); אין שטח גזרה ואין עלות תפירה', 'חסר נתוני עלות')
add(CAT_W, 'כובע + כפפה (סט)', 'N.B', 'פלנל/טריקו', None, None, '—', None,
    'עומר ביקש ~100 סטים שמנת + לבן; אין פריט/עלות במערכת', 'חסר נתוני עלות')

# ---------------------------------------------------------------- קיץ — טריקו לבן
CAT_S = 'קיץ — טריקו לבן'
for size in ['0-3', '3-6', '6-12', '12-18', '18-24', '24-30']:
    c, d = calc('בגד גוף ש.א', size, 'טריקו', ticks=TICKS['בגד גוף ש.א'][size], ribbon=1)
    add(CAT_S, 'בגד גוף טריקו שרוול ארוך', size, 'טריקו', c, r1(c * MARGIN), d, RIVHIT_SALE['בגד גוף טריקו ארוך'],
        'במחירון יש רק גופייה וקצר בטריקו')

c, d = calc('בגד גוף מעטפת', '0-3', 'טריקו', ticks=5, ribbon=1)
add(CAT_S, 'בגד גוף מעטפת טריקו', '0-3', 'טריקו', c, r1(c * MARGIN), d, RIVHIT_SALE['בגד גוף מעטפת טריקו'],
    'עומר ביקש מעטפת+רגלית טריקו 0-3')
for size in ['0-3', '3-6']:
    c, d = calc('רגליות', size, 'טריקו', elastic=1)
    add(CAT_S, 'רגליות טריקו', size, 'טריקו', c, r1(c * MARGIN), d,
        RIVHIT_SALE['רגליות טריקו' if size == '0-3' else 'רגליות טריקו 3-6'])
for size, area in (('0-3', 0.195), ('3-6', SQM['גטקס']['3-6'])):
    c, d = calc('גטקס', size, 'טריקו', elastic=1, sqm=area)
    add(CAT_S, 'גטקס פתוח טריקו', size, 'טריקו', c, r1(c * MARGIN), d, RIVHIT_SALE['גטקס טריקו'],
        'שטח 0-3 משוער כמו רגליות 0-3' if size == '0-3' else '', 'הערכה' if size == '0-3' else 'מחושב')

# גופיות סבא 2 / 12 / 14 / 16 — המשך סדרת המחירון לפי שטח גזרה (3 → 5.5 ... 10 → 6.4)
slope = (EXCEL_SABA['10'] - EXCEL_SABA['3']) / (SQM['גופיות סבא']['10'] - SQM['גופיות סבא']['3'])
for size, area in (('2', 0.150), ('12', SQM['גופיות סבא']['12']), ('14', SQM['גופיות סבא']['14']), ('16', SQM['גופיות סבא']['16'])):
    p = r1(EXCEL_SABA['10'] + (area - SQM['גופיות סבא']['10']) * slope)
    add(CAT_S, 'גופיות סבא', size, 'טריקו', None, p,
        f'המשך סדרת המחירון לפי שטח גזרה ({area:.3f} מ"ר, {slope:.2f} ₪/מ"ר)', RIVHIT_SALE.get(f'גופיות סבא {size}'),
        'שטח מידה 2 משוער' if size == '2' else '', 'הערכה' if size == '2' else 'מחושב')

# אוברול טריקו 0-3 — כמו אוברול פלנל 0-3 (12.5) בהפרש מחיר הבד
fab_diff = SQM['אוברול']['0-3'] * (FABRIC_PRICE['פלנל'] - FABRIC_PRICE['טריקו'])
add(CAT_S, 'אוברול טריקו (רוכסן)', '0-3', 'טריקו', None, r1(EXCEL_OVERALL['0-3'] - fab_diff * MARGIN),
    f'אוברול פלנל 0-3 במחירון (12.5) פחות הפרש בד {fab_diff:.2f}×1.2', RIVHIT_SALE['אוברול טריקו'], '', 'הערכה')

for name, key, note in (
    ('גופיות בייבי קצר טריקו', 'גופיות בייב קצר', 'מידות 0-4'),
    ('חזיות סגור טריקו N.B', 'חזיות סגור טריקו', ''),
    ('חזיות טיקטק טריקו + כפפה N.B', 'חזיות טיקטק טריקו', ''),
    ('תחתונים לבן', 'תחתונים', 'מידה 2'),
):
    s = RIVHIT_SALE[key]
    add(CAT_S, name, note.replace('מידות ', '').replace('מידה ', '') if note else 'N.B', 'טריקו', None,
        r1(s * STORE_RATIO), f'מחיר חנויות {s} × {STORE_RATIO}', s,
        'אין גזרה בקטלוג — לפי יחס ממוצע מחירון/חנויות', 'הערכה')

# סט חזיה + רגלית טריקו = סכום הרכיבים
mp = next(r for r in rows if r['name'] == 'בגד גוף מעטפת טריקו')['price']
rp = next(r for r in rows if r['name'] == 'רגליות טריקו' and r['size'] == '0-3')['price']
add(CAT_S, 'סט בגד גוף מעטפת + רגלית טריקו', '0-3', 'טריקו', None, r1(mp + rp),
    f'מעטפת טריקו {mp} + רגליות טריקו {rp}', RIVHIT_SALE['סט חזיה+רגלית טריקו'], 'הסט שעומר ביקש')

# ---------------------------------------------------------------- אחר
CAT_O = 'נלווים — לבן'
for name, key, note in (('חיתולי פלנל לבן (חבילה)', 'חיתולי פלנל', 'לבדוק כמות בחבילה'),
                        ('סדין + ציפה פלנל למיטת תינוק', 'סדין+ציפה פלנל', '')):
    s = RIVHIT_SALE[key]
    add(CAT_O, name, '', 'פלנל', None, r1(s * STORE_RATIO), f'מחיר חנויות {s} × {STORE_RATIO}', s, note, 'הערכה')


# ---------------------------------------------------------------- כבר מתומחר באפליקציה (לא במחירון האקסל)
def app_priced():
    try:
        from optitex_analyzer.core.data_processor import DataProcessor
        dp = DataProcessor()
        out = []
        for key, price in sorted((dp.baby_basic_price_list or {}).items()):
            if not str(key).startswith('nb|'):
                continue
            parts = str(key).split('|')[1:]
            out.append((parts[0], parts[1] if len(parts) > 1 else '', parts[2] if len(parts) > 2 else '', price))
        return out
    except Exception as exc:  # pragma: no cover
        print('warn: app price list not loaded:', exc)
        return []


# ---------------------------------------------------------------- Excel
HDR_FILL = PatternFill('solid', fgColor='1F3A5F')
HDR_FONT = Font(bold=True, color='FFFFFF', name='Arial', size=11)
CAT_FILL = PatternFill('solid', fgColor='E8EEF5')
EST_FILL = PatternFill('solid', fgColor='FFF4D6')
MISS_FILL = PatternFill('solid', fgColor='FADBD8')
THIN = Side(style='thin', color='B8C2CC')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BASE_FONT = Font(name='Arial', size=10)
BOLD = Font(name='Arial', size=10, bold=True)


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font, cell.border = HDR_FILL, HDR_FONT, BORDER
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def build():
    os.makedirs(OUT_DIR, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = 'השלמת מחירון לבן'
    ws.sheet_view.rightToLeft = True

    ws['A1'] = 'השלמת מחירון בייבי בייסיק 2026-27 — מוצרים לבנים שאינם במחירון'
    ws['A1'].font = Font(name='Arial', size=14, bold=True)
    ws['A2'] = (f'נוצר {datetime.now():%d/%m/%Y}. מחיר מוצע = עלות × 1.20 (כמו כלל "מכירה לבייבי בייסיק" '
                f'בטאב טורקיה). שורות צהובות = הערכה, אדומות = חסר נתוני עלות.')
    ws['A2'].font = Font(name='Arial', size=10, italic=True, color='555555')

    headers = ['קטגוריה', 'שם מוצר', 'מידה', 'בד', 'עלות מחושבת ₪', 'מחיר מוצע 2026-27 ₪',
               'מחיר חנויות אריה ₪', 'סטטוס', 'בסיס החישוב', 'הערות']
    ws.append([])
    ws.append(headers)
    style_header(ws, 4, len(headers))
    r = 5
    last_cat = None
    for row in rows:
        if row['cat'] != last_cat:
            ws.cell(row=r, column=1, value=row['cat']).font = BOLD
            for c in range(1, len(headers) + 1):
                ws.cell(row=r, column=c).fill = CAT_FILL
                ws.cell(row=r, column=c).border = BORDER
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(headers))
            r += 1
            last_cat = row['cat']
        vals = [row['cat'], row['name'], row['size'], row['fabric'],
                round(row['cost'], 2) if row['cost'] is not None else None,
                row['price'], row['store'], row['status'], row['basis'], row['note']]
        for c, v in enumerate(vals, start=1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = BOLD if c == 6 else BASE_FONT
            cell.border = BORDER
            cell.alignment = Alignment(horizontal='center' if c in (3, 4, 5, 6, 7, 8) else 'right',
                                       vertical='center', wrap_text=c in (9, 10))
            if c in (5, 6, 7) and isinstance(v, (int, float)):
                cell.number_format = '0.00'
            if row['status'] == 'הערכה':
                cell.fill = EST_FILL
            elif row['status'] == 'חסר נתוני עלות':
                cell.fill = MISS_FILL
        r += 1

    widths = [18, 36, 9, 9, 13, 16, 15, 14, 60, 44]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A5'

    # ---- sheet 2: method
    ws2 = wb.create_sheet('שיטת חישוב')
    ws2.sheet_view.rightToLeft = True
    lines = [
        ('מחיר מוצע', 'עלות × 1.20 — אותו כלל של "מכירה לבייבי בייסיק בעלות" בטאב "טורקיה מחירון" במחירון 2026-27.'),
        ('עלות', 'בד + טיקטקים + גומי + סרט + תפירה.'),
        ('בד', 'שטח גזרה (מ"ר) מתוך קטלוג המוצרים × מחיר למ"ר: פלנל לבן 11 ₪, טריקו לבן 10 ₪ (fabric_prices).'),
        ('נלווים', 'טיקטק 0.30 ₪ ליח\', גומי 0.50 ₪, סרט/תווית 0.50 ₪ (כמו ב"מחירי עלות אריה.xlsx").'),
        ('תפירה', 'בגד גוף ש.א 2.75 · מעטפת 3.10 · רגליות 1.75 · גטקס 1.95 · גופיות פלנל 3.10 · חולצת בייבי 2.75 · אוברול 4.10.'),
        ('אימות השיטה', 'חישוב כזה לבגד גוף טריקו קצר 0-3 נותן 7.3 מול 7.5 במחירון; גופייה טריקו 0-3 נותן 6.2 מול 6.5 — סטייה של עד ~5%.'),
        ('הערכה', 'פריטים ללא שטח גזרה בקטלוג: הושלמו לפי המשך סדרה או לפי יחס ממוצע 0.66 בין מחיר בייבי בייסיק למחיר החנויות של אריה בריווחית.'),
        ('חסר נתוני עלות', 'אוברול ארוך סואגי, כובע+כפפה — אין גזרה/תפירה במערכת; יש להוסיף לקטלוג המוצרים ולהריץ שוב.'),
        ('הרצה מחדש', 'python scripts/build_baby_basic_missing_prices.py'),
    ]
    ws2.append(['נושא', 'פירוט'])
    style_header(ws2, 1, 2)
    for k, v in lines:
        ws2.append([k, v])
    ws2.column_dimensions['A'].width = 18
    ws2.column_dimensions['B'].width = 120
    for row in ws2.iter_rows(min_row=2):
        for cell in row:
            cell.font = BASE_FONT
            cell.alignment = Alignment(horizontal='right', vertical='top', wrap_text=True)
            cell.border = BORDER

    # ---- sheet 3: already priced in app but missing from the Excel
    ws3 = wb.create_sheet('כבר מתומחר באפליקציה')
    ws3.sheet_view.rightToLeft = True
    ws3.append(['דגם', 'בד', 'צבע', 'מחיר באפליקציה ₪', 'הערה'])
    style_header(ws3, 1, 5)
    for model, fabric, color, price in app_priced():
        ws3.append([model, fabric, color, price, 'קיים במחירון האפליקציה (ללא ברקוד) אך לא במחירון האקסל'])
    for i, w in enumerate([28, 10, 10, 18, 60], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    for row in ws3.iter_rows(min_row=2):
        for cell in row:
            cell.font = BASE_FONT
            cell.border = BORDER
            cell.alignment = Alignment(horizontal='right', vertical='center')

    wb.save(OUT_FILE)
    return OUT_FILE


def attach_to_thread(path: str):
    try:
        from optitex_analyzer.core.data_processor import DataProcessor
        dp = DataProcessor()
        thread = dp.get_baby_basic_email_thread(PRICE_LIST_THREAD_ID)
        if not thread:
            return False
        base = os.path.basename(path)
        for att in thread.get('extra_attachments', []) or []:
            if os.path.basename(att.get('path', '')) == base:
                # refresh the stored copy
                import shutil
                shutil.copy2(path, att['path'])
                return True
        return bool(dp.add_baby_basic_attachment(PRICE_LIST_THREAD_ID, path, label='השלמת מחירון לבן 2026-27 (מחושב)'))
    except Exception as exc:  # pragma: no cover
        print('warn: attach failed:', exc)
        return False


if __name__ == '__main__':
    out = build()
    attached = attach_to_thread(out)
    calc_n = sum(1 for r in rows if r['status'] == 'מחושב')
    est_n = sum(1 for r in rows if r['status'] == 'הערכה')
    miss_n = sum(1 for r in rows if r['status'] == 'חסר נתוני עלות')
    print(f'saved: {out}')
    print(f'rows={len(rows)} calculated={calc_n} estimated={est_n} missing={miss_n} attached={attached}')
    for r_ in rows:
        print(f"{r_['status']:<14} | {r_['name']} {r_['size']} | cost={r_['cost'] if r_['cost'] is None else round(r_['cost'], 2)} | price={r_['price']}")
