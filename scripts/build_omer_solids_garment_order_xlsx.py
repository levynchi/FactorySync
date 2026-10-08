# -*- coding: utf-8 -*-
"""Sendable Excel order for Omer — solid interlock, 4 products."""
from __future__ import annotations

import json
import shutil
import sys
from datetime import datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from PIL import Image as PILImage

from omer_solids_order_data import (
    BABY_QTY,
    BABY_SIZES,
    COLORS,
    DATE,
    FABRIC,
    NEWBORN_QTY,
    OUT_XLSX_NAME,
    OVERALL_QTY,
    PAJAMA_QTY,
    PAJAMA_SIZES,
    PRODUCTS,
    totals,
)

ROOT = Path(__file__).resolve().parents[1]
OUT_XLSX = ROOT / "exports" / OUT_XLSX_NAME
SUPPLIER_DIR = ROOT / "supplier_documents" / "4"
REFS_DIR = SUPPLIER_DIR / "order_refs"
SUPPLIERS_JSON = ROOT / "suppliers.json"
ASSETS = Path(r"C:\Users\levyn\.cursor\projects\c-optitex-excell\assets")

LOOK_SOURCES = {
    "overall": "c__Users_levyn_AppData_Roaming_Cursor_User_workspaceStorage_89e54074c5d09557f2fae04a3874ff96_images_______-62e4574f-a859-49cc-858a-c794f14630b2.jpg",
    "pajama": "c__Users_levyn_AppData_Roaming_Cursor_User_workspaceStorage_89e54074c5d09557f2fae04a3874ff96_images_pigama_set-143a7a7f-7cd5-4cf8-b209-cb0843b988b5.jpg",
    "baby": "c__Users_levyn_AppData_Roaming_Cursor_User_workspaceStorage_89e54074c5d09557f2fae04a3874ff96_images_baby_set-2ae3e8fc-6cda-4aad-a6b9-58580f4bceb0.jpg",
    "newborn": "c__Users_levyn_AppData_Roaming_Cursor_User_workspaceStorage_89e54074c5d09557f2fae04a3874ff96_images_new_born_set-20fd5e3f-5380-4999-a4ef-5fcd6b01be06.jpg",
}
LOOK_FILES = {
    "overall": REFS_DIR / "overall_look.jpg",
    "pajama": REFS_DIR / "pajama_set_look.jpg",
    "baby": REFS_DIR / "baby_set_look.jpg",
    "newborn": REFS_DIR / "newborn_set_look.jpg",
}
THUMB_FILES = {
    key: REFS_DIR / f"{path.stem}_thumb.jpg" for key, path in LOOK_FILES.items()
}

NAVY = PatternFill("solid", fgColor="1B2A4A")
GOLD = PatternFill("solid", fgColor="B0894A")
ALT = PatternFill("solid", fgColor="F3EFE8")
PAPER = PatternFill("solid", fgColor="FAF7F2")
AMBER = PatternFill("solid", fgColor="FEF3C7")
TOTAL = PatternFill("solid", fgColor="1B2A4A")
SECTION = PatternFill("solid", fgColor="E8EEF6")
WHITE = PatternFill("solid", fgColor="FFFFFF")
THIN = Border(
    left=Side(style="thin", color="D6D3D1"),
    right=Side(style="thin", color="D6D3D1"),
    top=Side(style="thin", color="D6D3D1"),
    bottom=Side(style="thin", color="D6D3D1"),
)
TITLE = Font(name="Calibri", bold=True, size=16, color="1B2A4A")
SUB = Font(name="Calibri", size=11, color="57534E")
HEAD = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
BODY = Font(name="Calibri", size=10, color="1C1917")
BOLD = Font(name="Calibri", bold=True, size=10, color="1C1917")
WHITE_BOLD = Font(name="Calibri", bold=True, size=10, color="FFFFFF")
SECTION_FONT = Font(name="Calibri", bold=True, size=11, color="1B2A4A")
NOTE = Font(name="Calibri", italic=True, size=9, color="57534E")
AMBER_FONT = Font(name="Calibri", size=10, color="92400E")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
HEADERS = ["#", "Product", "Models / pieces", "Size", "Color", "Akdem code", "Qty", "Unit", "Pattern status"]
NCOLS = len(HEADERS)
LOOK_COL = 11  # column K — keep a gap after the table
LOOK_MAX_W = 168
LOOK_MAX_H = 210


def _save_fit(src: Path, dest: Path, box: tuple[int, int], quality: int = 88) -> Path:
    with PILImage.open(src) as im:
        im = im.convert("RGB")
        im.thumbnail(box, PILImage.Resampling.LANCZOS)
        dest.parent.mkdir(parents=True, exist_ok=True)
        im.save(dest, "JPEG", quality=quality, optimize=True)
    return dest


def prepare_look_images() -> None:
    REFS_DIR.mkdir(parents=True, exist_ok=True)
    for key, src_name in LOOK_SOURCES.items():
        src = ASSETS / src_name
        if not src.exists():
            raise FileNotFoundError(f"Missing look image: {src}")
        _save_fit(src, LOOK_FILES[key], (240, 300))
        thumb_box = (LOOK_MAX_W, 155) if key == "newborn" else (LOOK_MAX_W, LOOK_MAX_H)
        _save_fit(src, THUMB_FILES[key], thumb_box)


def _add_look(ws: Worksheet, path: Path, row: int) -> None:
    img = XLImage(str(path))
    with PILImage.open(path) as im:
        img.width, img.height = im.size
    img.anchor = f"K{row}"
    ws.add_image(img)


def _paint(cell, *, fill=None, font=None, align=None, border=True) -> None:
    if fill is not None:
        cell.fill = fill
    if font is not None:
        cell.font = font
    if align is not None:
        cell.alignment = align
    if border:
        cell.border = THIN


def _merge_row(ws: Worksheet, row: int, value: str, font, fill, align=LEFT) -> None:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=NCOLS)
    cell = ws.cell(row=row, column=1, value=value)
    _paint(cell, fill=fill, font=font, align=align, border=False)
    for col in range(2, NCOLS + 1):
        _paint(ws.cell(row=row, column=col), fill=fill, border=False)


def _header_row(ws: Worksheet, row: int) -> None:
    for col, title in enumerate(HEADERS, 1):
        cell = ws.cell(row=row, column=col, value=title)
        _paint(cell, fill=NAVY, font=HEAD, align=CENTER)
    ws.row_dimensions[row].height = 22
    ws.auto_filter.ref = f"A{row}:I{row}"
    ws.freeze_panes = f"A{row + 1}"


def _line(
    ws: Worksheet,
    row: int,
    values: list,
    *,
    total: bool = False,
    alt: bool = False,
) -> int:
    fill = TOTAL if total else (ALT if alt else WHITE)
    font = WHITE_BOLD if total else BODY
    for col, val in enumerate(values, 1):
        cell = ws.cell(row=row, column=col, value=val)
        align = CENTER if col in (1, 4, 6, 7, 8) else LEFT
        _paint(cell, fill=fill, font=font, align=align)
        if col == 7 and isinstance(val, int):
            cell.number_format = "#,##0"
    ws.row_dimensions[row].height = 20
    return row + 1


def _section(ws: Worksheet, row: int, text: str) -> int:
    _merge_row(ws, row, text, SECTION_FONT, SECTION)
    ws.row_dimensions[row].height = 22
    return row + 1


def build_order_sheet(ws: Worksheet) -> None:
    t = totals()
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_title_rows = "10:10"
    ws.oddHeader.left.text = "Aryeh Baby Clothes  ·  Levi Moshe"
    ws.oddHeader.right.text = "Omer / Nibby Baby"
    ws.oddFooter.left.text = "Confidential — production order"
    ws.oddFooter.right.text = "Page &P of &N"

    widths = [5, 22, 32, 12, 16, 14, 10, 8, 48]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.column_dimensions["J"].width = 3
    ws.column_dimensions["K"].width = 26
    look_cell = ws.cell(row=9, column=LOOK_COL, value="Look")
    _paint(look_cell, fill=NAVY, font=HEAD, align=CENTER)

    _merge_row(ws, 1, "GARMENT PRODUCTION ORDER — SOLID INTERLOCK", TITLE, PAPER)
    ws.row_dimensions[1].height = 26
    _merge_row(
        ws,
        2,
        f"From Lior  ·  Aryeh Baby Clothes / Levi Moshe    |    To Omer  ·  Nibby Baby  ·  Turkey    |    {DATE}",
        SUB,
        PAPER,
    )
    _merge_row(
        ws,
        3,
        "4 products  ·  4 colors  ·  FOB USD quote requested (excluding Turkish VAT)",
        SUB,
        PAPER,
    )

    fabric_line = "  ·  ".join(f"{k}: {v}" for k, v in FABRIC)
    _merge_row(ws, 5, f"FABRIC  —  {fabric_line}", BOLD, PAPER)
    colors_line = "   |   ".join(
        f"{c['name']}  {c['code']}" + (f"  ({c['note']})" if c["key"] == "offwhite" else "")
        for c in COLORS
    )
    _merge_row(ws, 6, f"COLORS  —  {colors_line}", BODY, PAPER)
    _merge_row(
        ws,
        7,
        "Fabric source: Bilal / Akdem not available lately for these solids. Another mill is OK if spec and colors match. "
        "Keep leftover fabric at your place; I will order from that stock. Remaining leftover I will buy from you.",
        NOTE,
        PAPER,
    )
    ws.row_dimensions[7].height = 32

    row = 9
    _header_row(ws, row)
    row = 10
    alt = False

    # A Overall
    overall_look_row = row
    row = _section(ws, row, "A  ·  OVERALL / FOOTED ROMPER  ·  closed feet, double zipper  ·  NEW pattern (not sewn yet)")
    _add_look(ws, THUMB_FILES["overall"], overall_look_row)
    for size, qty in OVERALL_QTY.items():
        for col in COLORS:
            row = _line(
                ws,
                row,
                ["A", "Overall", "New overall / sleepsuit", size, col["name"], col["code"], qty, "pcs", PRODUCTS[0]["pattern"]],
                alt=alt,
            )
            alt = not alt
    row = _line(ws, row, ["A", "Overall total", "", "", "", "", t["overall"], "pcs", ""], total=True)

    # B Pajama
    pajama_look_row = row
    row = _section(ws, row, "B  ·  PAJAMA SET  ·  new long-sleeve shirt + model 202 open gatkes  ·  202 already sewn, shirt is new")
    _add_look(ws, THUMB_FILES["pajama"], pajama_look_row)
    for size in PAJAMA_SIZES:
        for col in COLORS:
            row = _line(
                ws,
                row,
                ["B", "Pajama set", "New shirt + 202 open pants", size, col["name"], col["code"], PAJAMA_QTY, "sets", PRODUCTS[1]["pattern"]],
                alt=alt,
            )
            alt = not alt
    row = _line(ws, row, ["B", "Pajama set total", "", "", "", "", t["pajama"], "sets", ""], total=True)

    # C Baby set
    baby_look_row = row
    row = _section(ws, row, "C  ·  BABY SET  ·  201 envelope wrap bodysuit + footed pants  ·  both patterns already sewn  ·  3-6 / 6-12")
    _add_look(ws, THUMB_FILES["baby"], baby_look_row)
    for size in BABY_SIZES:
        for col in COLORS:
            row = _line(
                ws,
                row,
                ["C", "Baby set", "201 wrap + footed pants", size, col["name"], col["code"], BABY_QTY, "sets", PRODUCTS[2]["pattern"]],
                alt=alt,
            )
            alt = not alt
    row = _line(ws, row, ["C", "Baby set total", "", "", "", "", t["baby"], "sets", ""], total=True)

    # D Newborn last (0-3)
    newborn_look_row = row
    row = _section(ws, row, "D  ·  NEWBORN SET  ·  201 envelope wrap bodysuit + footed pants  ·  both patterns already sewn  ·  0-3 only")
    _add_look(ws, THUMB_FILES["newborn"], newborn_look_row)
    for col in COLORS:
        row = _line(
            ws,
            row,
            ["D", "Newborn set", "201 wrap + footed pants", "0-3", col["name"], col["code"], NEWBORN_QTY[col["key"]], "sets", PRODUCTS[3]["pattern"]],
            alt=alt,
        )
        alt = not alt
    row = _line(ws, row, ["D", "Newborn set total", "", "0-3", "", "", t["newborn"], "sets", ""], total=True)

    row += 2
    row = _section(ws, row, "ORDER TOTALS")
    _header_row(ws, row)
    # don't reset freeze - already set. Re-applying header style on totals is ok but auto_filter would move.
    # Revert auto_filter to first header
    ws.auto_filter.ref = "A9:I9"
    ws.freeze_panes = "A10"
    row += 1
    summary = [
        ["A", "Overall / footed romper", "New overall", "0-3 / 3-6 / 6-12", "", "", t["overall"], "pcs", PRODUCTS[0]["pattern"]],
        ["B", "Pajama set", "New shirt + 202", "12-18 / 18-24 / 24-30", "", "", t["pajama"], "sets", PRODUCTS[1]["pattern"]],
        ["C", "Baby set", "201 + footed pants", "3-6 / 6-12", "", "", t["baby"], "sets", PRODUCTS[2]["pattern"]],
        ["D", "Newborn set", "201 + footed pants", "0-3", "", "", t["newborn"], "sets", PRODUCTS[3]["pattern"]],
        ["", "GRAND TOTAL", "", "", "", "", t["units"], "", "A pcs + B/C/D sets"],
    ]
    for i, vals in enumerate(summary):
        row = _line(ws, row, vals, total=i == len(summary) - 1, alt=i % 2 == 1)

    row += 1
    _merge_row(
        ws,
        row,
        "PLEASE QUOTE: FOB USD per piece (A) and per set (B, C, D), excluding Turkish VAT. "
        "Use the markers for fabric consumption. Include cutting, sewing, zipper / snaps, labels, packing. "
        "Please confirm brand / labels (Arye or Baby Basic) before production.",
        AMBER_FONT,
        AMBER,
    )
    ws.row_dimensions[row].height = 36
    row += 2
    _merge_row(ws, row, "Thank you, Lior", BOLD, PAPER)
    ws.print_area = f"A1:K{row}"


def build_notes_sheet(ws: Worksheet) -> None:
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 92
    ws.merge_cells("A1:B1")
    ws["A1"] = "NOTES FOR THIS ORDER"
    ws["A1"].font = TITLE
    ws["A1"].fill = PAPER
    ws["B1"].fill = PAPER
    ws.row_dimensions[1].height = 24

    rows = [
        ("Fabric", "100% cotton interlock, 30/1, 220 GSM, brushed one side. Solid colors only."),
        ("Mill", "Bilal / Akdem not available lately. Another mill is OK if spec and colors match the codes / simulations."),
        ("Leftover fabric", "Keep leftover at your place. I will order from that stock from time to time. Remaining leftover I will buy from you."),
        ("A Overall", "New cut. You have not sewn this overall / footed romper yet. Closed feet, double zipper."),
        ("B Pajama set", "Model 202 open gatkes — you already sewed. Long-sleeve shirt — new pattern, not sewn yet."),
        ("C Baby set", "201 envelope wrap bodysuit + footed pants. You already sewed both patterns. Sizes 3-6 and 6-12."),
        ("D Newborn set", "Same as Baby set (201 + footed pants). Size 0-3 only. Gray 160; other colors 120 each."),
        ("Quote", "FOB USD, excluding Turkish VAT. Per piece for A, per set for B / C / D."),
        ("Brand / labels", "Please confirm Arye or Baby Basic before production. Quantities are not split by brand yet."),
    ]
    r = 3
    for label, value in rows:
        a = ws.cell(row=r, column=1, value=label)
        b = ws.cell(row=r, column=2, value=value)
        _paint(a, fill=NAVY, font=HEAD, align=LEFT)
        _paint(b, fill=PAPER, font=BODY, align=LEFT)
        ws.row_dimensions[r].height = 28
        r += 1


def build_looks_sheet(ws: Worksheet) -> None:
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.merge_cells("A1:D1")
    ws["A1"] = "REFERENCE LOOKS  ·  solid interlock order"
    ws["A1"].font = TITLE
    ws["A1"].fill = PAPER
    for col in range(1, 5):
        ws.cell(row=1, column=col).fill = PAPER
    ws.row_dimensions[1].height = 26
    ws.merge_cells("A2:D2")
    ws["A2"] = "Illustration only — colors in the photos are examples. Produce in Blue Navy, Pink, Gray, Off White as per the Order sheet."
    ws["A2"].font = NOTE
    ws["A2"].alignment = LEFT
    ws.row_dimensions[2].height = 20

    blocks = [
        (3, 1, "A  ·  Overall", LOOK_FILES["overall"]),
        (3, 3, "B  ·  Pajama set", LOOK_FILES["pajama"]),
        (20, 1, "C  ·  Baby set", LOOK_FILES["baby"]),
        (20, 3, "D  ·  Newborn set", LOOK_FILES["newborn"]),
    ]
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 28
    for row, col, title, path in blocks:
        cell = ws.cell(row=row, column=col, value=title)
        _paint(cell, fill=NAVY, font=HEAD, align=LEFT)
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        ws.cell(row=row, column=col + 1).fill = NAVY
        img = XLImage(str(path))
        with PILImage.open(path) as im:
            img.width, img.height = im.size
        img.anchor = f"{get_column_letter(col)}{row + 1}"
        ws.add_image(img)
        ws.row_dimensions[row].height = 20


def register_document(src: Path, filename: str) -> Path:
    SUPPLIER_DIR.mkdir(parents=True, exist_ok=True)
    dest = SUPPLIER_DIR / filename
    shutil.copy2(src, dest)
    with open(SUPPLIERS_JSON, "r", encoding="utf-8") as f:
        suppliers = json.load(f)
    omer = next(s for s in suppliers if s.get("id") == 4)
    docs = omer.setdefault("documents", [])
    if not any(d.get("filename") == filename for d in docs):
        next_id = max((d.get("id", 0) for d in docs), default=0) + 1
        docs.append({
            "id": next_id,
            "filename": filename,
            "original_name": filename,
            "uploaded_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        })
        with open(SUPPLIERS_JSON, "w", encoding="utf-8") as f:
            json.dump(suppliers, f, ensure_ascii=False, indent=2)
            f.write("\n")
    return dest


def main() -> None:
    OUT_XLSX.parent.mkdir(parents=True, exist_ok=True)
    prepare_look_images()
    wb = Workbook()
    order = wb.active
    order.title = "Order"
    build_order_sheet(order)
    looks = wb.create_sheet("Looks")
    build_looks_sheet(looks)
    notes = wb.create_sheet("Notes")
    build_notes_sheet(notes)
    wb.save(OUT_XLSX)
    dest = register_document(OUT_XLSX, OUT_XLSX_NAME)
    print(f"Wrote {OUT_XLSX}")
    print(f"Copied {dest}")


if __name__ == "__main__":
    main()
