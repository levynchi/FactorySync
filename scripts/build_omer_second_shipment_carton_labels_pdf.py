# -*- coding: utf-8 -*-
"""Carton labels (105x99 mm, 6 per A4 page) for Omer second shipment 025/204.

Reads supplier_documents/4/025-204 second shipment_*.xlsx (sheet Sayfa1),
writes one PDF per brand (Arye / Baby Basic), 4 identical labels per carton.
Each label: logo, carton number, brand, model, size, flannel / not flannel,
pieces + packets.
"""
from __future__ import annotations

import os
import re
import sys
from pathlib import Path

from openpyxl import load_workbook
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from optitex_analyzer.core.label_sheet import (  # noqa: E402
    LOGO_BABY_BASIC_PATH,
    LOGO_PATH,
    _FONT_BOLD,
    _FONT_LIGHT,
    _FONT_NAME,
    _draw_image_fit,
    _get_logo,
    _register_fonts,
    _shape,
)

SUPPLIER_DIR = ROOT / "supplier_documents" / "4"
PACKING_PREFIX = "025-204 second shipment"

# ---- geometry: 2 x 3 labels of 105 x 99 mm on A4 (210 x 297) ----
PAGE_W, PAGE_H = A4
LABEL_W = 105 * mm
LABEL_H = 99 * mm
COLS = 2
ROWS = 3
PER_PAGE = COLS * ROWS
BORDER_INSET = 2 * mm      # dashed frame inset from cell edge
CONTENT_PAD = 5 * mm       # inner padding between frame and content
LOGO_H = 16 * mm
COPIES_PER_CARTON = 4

MODEL_NAMES = {
    25: "בגד גוף שרוול ארוך",
    204: "גופיות פלנל",
}
FLANNEL_MODELS = {25, 204}

OUTPUTS = {
    "ARYE": {
        "dest": SUPPLIER_DIR / "מדבקות_קרטונים_105x99_אריה_025-204.pdf",
        "logo": LOGO_PATH,
    },
    "BABY BASIC": {
        "dest": SUPPLIER_DIR / "מדבקות_קרטונים_105x99_בייבי_בייסיק_025-204.pdf",
        "logo": LOGO_BABY_BASIC_PATH,
    },
}


# ------------------------------------------------------------------ parsing
def _find_packing_list() -> Path:
    for name in sorted(os.listdir(SUPPLIER_DIR)):
        if name.startswith("~$"):
            continue
        if name.startswith(PACKING_PREFIX) and name.lower().endswith(".xlsx"):
            return SUPPLIER_DIR / name
    raise FileNotFoundError(f"{PACKING_PREFIX}*.xlsx not found in {SUPPLIER_DIR}")


def _normalize_brand(raw) -> str:
    text = str(raw or "").strip().upper().replace("İ", "I").replace("ı", "I")
    if "BABY" in text:
        return "BABY BASIC"
    if "ARYE" in text:
        return "ARYE"
    return text


def _parse_int(raw) -> int | None:
    if raw is None:
        return None
    if isinstance(raw, (int, float)):
        return int(raw)
    digits = re.sub(r"\D", "", str(raw))
    return int(digits) if digits else None


def load_cartons(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    cartons: list[dict] = []
    header_seen = False
    for row in ws.iter_rows(values_only=True):
        if not row or all(c is None or str(c).strip() == "" for c in row[:6]):
            continue
        first = str(row[0] or "").strip().upper()
        if not header_seen:
            if "CARTON" in first:
                header_seen = True
            continue
        if row[1] is None or str(row[1]).strip() == "":
            continue  # subtotal row
        carton_label = str(row[0] or "").strip().upper()
        carton_no = _parse_int(carton_label)
        model = _parse_int(row[1])
        pieces = _parse_int(row[5])
        if carton_no is None or model is None or pieces is None:
            continue
        cartons.append({
            "carton": carton_no,
            "carton_label": carton_label,
            "model": model,
            "size": str(row[2] or "").strip(),
            "brand": _normalize_brand(row[3]),
            "packets": _parse_int(row[4]) or 0,
            "pieces": pieces,
        })
    if not cartons:
        raise RuntimeError(f"No carton rows found in {path.name}")
    return cartons


# ------------------------------------------------------------------ drawing
def _fit_size(c: canvas.Canvas, text: str, font: str, max_size: float, min_size: float, max_width: float) -> float:
    size = max_size
    while size > min_size and c.stringWidth(text, font, size) > max_width:
        size -= 0.5
    return size


def _draw_qty_line(c, mid_x, y, font, size, pieces, packets, max_w) -> None:
    """Units on the right, packets on the left (keeps Hebrew/number order)."""
    units = _shape(f"{pieces} יחידות")
    packs = _shape(f"{packets} מארזים")
    sep = "·"
    pad = 8
    c.setFont(font, size)
    w_units = c.stringWidth(units, font, size)
    w_packs = c.stringWidth(packs, font, size)
    w_sep = c.stringWidth(sep, font, size)
    total = w_units + w_packs + w_sep + pad * 2
    while size > 9 and total > max_w:
        size -= 0.5
        c.setFont(font, size)
        w_units = c.stringWidth(units, font, size)
        w_packs = c.stringWidth(packs, font, size)
        w_sep = c.stringWidth(sep, font, size)
        total = w_units + w_packs + w_sep + pad * 2
    right = mid_x + total / 2.0
    c.drawRightString(right, y, units)
    c.drawCentredString(right - w_units - pad - w_sep / 2.0, y, sep)
    c.drawRightString(right - w_units - pad * 2 - w_sep, y, packs)


def _draw_fabric_badge(c: canvas.Canvas, mid_x: float, y_center: float, text: str, size: float) -> None:
    shaped = _shape(text)
    c.setFont(_FONT_BOLD, size)
    tw = c.stringWidth(shaped, _FONT_BOLD, size)
    pad_x, pad_y = 4 * mm, 1.6 * mm
    bw, bh = tw + 2 * pad_x, size + 2 * pad_y
    c.saveState()
    c.setLineWidth(1.0)
    c.setDash()
    c.roundRect(mid_x - bw / 2.0, y_center - bh / 2.0, bw, bh, 2.5 * mm, stroke=1, fill=0)
    c.restoreState()
    c.setFont(_FONT_BOLD, size)
    c.drawCentredString(mid_x, y_center - size * 0.35, shaped)


def _draw_carton_label(c: canvas.Canvas, x_left: float, y_bottom: float, item: dict, logo) -> None:
    # dashed frame
    c.saveState()
    c.setLineWidth(0.4)
    c.setDash(3, 3)
    c.setStrokeGray(0.55)
    c.roundRect(x_left + BORDER_INSET, y_bottom + BORDER_INSET,
                LABEL_W - 2 * BORDER_INSET, LABEL_H - 2 * BORDER_INSET, 2 * mm, stroke=1, fill=0)
    c.restoreState()

    x0 = x_left + BORDER_INSET + CONTENT_PAD
    x1 = x_left + LABEL_W - BORDER_INSET - CONTENT_PAD
    y0 = y_bottom + BORDER_INSET + CONTENT_PAD
    y1 = y_bottom + LABEL_H - BORDER_INSET - CONTENT_PAD
    mid_x = (x0 + x1) / 2.0
    content_w = x1 - x0

    if logo is not None:
        _draw_image_fit(c, logo, x0, y1 - LOGO_H, content_w, LOGO_H, anchor="n")
        y1 -= LOGO_H + 2 * mm

    model = int(item["model"])
    carton_text = item["carton_label"]
    model_text = MODEL_NAMES.get(model) or str(model)
    size = item["size"]
    size_text = f"מידה {size}"
    fabric_text = "פלנל" if model in FLANNEL_MODELS else "לא פלנל"
    pieces = int(item["pieces"])
    packets = int(item["packets"])

    number_size = 72
    model_size = _fit_size(c, _shape(model_text), _FONT_BOLD, 28, 14, content_w)
    size_size = _fit_size(c, _shape(size_text), _FONT_BOLD, 26, 14, content_w)
    badge_size = 21
    badge_h = badge_size + 2 * 1.6 * mm
    qty_size = 16
    gap = 6
    block_h = (number_size * 0.78 + model_size + size_size + badge_h + qty_size
               + gap * 4)
    avail = y1 - y0
    y = y1 - max(0.0, (avail - block_h) / 2.0)

    # carton number
    y -= number_size * 0.78
    c.setFont(_FONT_BOLD, number_size)
    c.drawCentredString(mid_x, y, carton_text)
    # model
    y -= gap + model_size
    c.setFont(_FONT_BOLD, model_size)
    c.drawCentredString(mid_x, y, _shape(model_text))
    # size
    y -= gap + size_size
    c.setFont(_FONT_BOLD, size_size)
    c.drawCentredString(mid_x, y, _shape(size_text))
    # fabric badge
    y -= gap + badge_h
    _draw_fabric_badge(c, mid_x, y + badge_h / 2.0, fabric_text, badge_size)
    # quantities
    y -= gap + qty_size
    _draw_qty_line(c, mid_x, y, _FONT_LIGHT, qty_size, pieces, packets, content_w)


def _cell_origin(index: int) -> tuple[float, float]:
    pos = index % PER_PAGE
    row = pos // COLS
    col = pos % COLS
    left_margin = (PAGE_W - COLS * LABEL_W) / 2.0
    top_margin = (PAGE_H - ROWS * LABEL_H) / 2.0
    x_left = left_margin + col * LABEL_W
    y_bottom = PAGE_H - top_margin - (row + 1) * LABEL_H
    return x_left, y_bottom


def build_brand_pdf(cartons: list[dict], dest: Path, logo_path: str) -> int:
    items: list[dict] = []
    for carton in sorted(cartons, key=lambda c: int(c["carton"])):
        items.extend([carton] * COPIES_PER_CARTON)
    if not items:
        raise ValueError(f"אין קרטונים ל-{dest.name}")

    _register_fonts()
    logo = _get_logo(logo_path)
    if logo is None:
        raise FileNotFoundError(f"לוגו לא נמצא: {logo_path}")
    dest.parent.mkdir(parents=True, exist_ok=True)
    tmp = dest.with_name(dest.stem + ".__tmp.pdf")
    c = canvas.Canvas(str(tmp), pagesize=(PAGE_W, PAGE_H))
    c.setTitle(dest.stem)
    for idx, item in enumerate(items):
        if idx > 0 and idx % PER_PAGE == 0:
            c.showPage()
        x_left, y_bottom = _cell_origin(idx)
        _draw_carton_label(c, x_left, y_bottom, item, logo)
    c.showPage()
    c.save()
    try:
        os.replace(tmp, dest)
    except PermissionError:
        alt = dest.with_name(dest.stem + "_מעודכן.pdf")
        os.replace(tmp, alt)
        print(f"הקובץ פתוח, נשמר במקום: {alt.name}")
    return len(items)


def main() -> None:
    packing = _find_packing_list()
    cartons = load_cartons(packing)
    by_brand: dict[str, list[dict]] = {key: [] for key in OUTPUTS}
    unknown = [c for c in cartons if c["brand"] not in by_brand]
    if unknown:
        raise RuntimeError(f"מותגים לא מזוהים: {sorted({c['brand'] for c in unknown})}")
    for item in cartons:
        by_brand[item["brand"]].append(item)

    print(f"source: {packing.name} cartons={len(cartons)}")
    for brand, spec in OUTPUTS.items():
        count = build_brand_pdf(by_brand[brand], spec["dest"], spec["logo"])
        pages = (count + PER_PAGE - 1) // PER_PAGE
        print(f"{spec['dest'].name}: cartons={len(by_brand[brand])} labels={count} pages={pages}")


if __name__ == "__main__":
    main()
