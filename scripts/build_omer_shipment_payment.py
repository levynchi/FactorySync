# -*- coding: utf-8 -*-
"""Build USD payment Excel for Omer first shipment (models 200/201/202).

Reads LEO 2026 price list + packing list 200-201-202.xlsx, writes a 3-sheet
workbook under supplier_documents/4/, and registers it in suppliers.json.
Prices are USD excluding Turkish VAT.
"""
from __future__ import annotations

import json
import os
import re
import unicodedata
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SUPPLIER_DIR = ROOT / "supplier_documents" / "4"
OUT_NAME = "omer first shipment payment USD 07.09.26.xlsx"
OUT_PATH = SUPPLIER_DIR / OUT_NAME
SUPPLIERS_JSON = ROOT / "suppliers.json"

HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
HEADER_FONT = Font(name="Segoe UI", bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(name="Segoe UI", bold=True, size=14, color="1E293B")
SUB_FONT = Font(name="Segoe UI", size=10, color="475569")
BODY_FONT = Font(name="Segoe UI", size=10)
BOLD_FONT = Font(name="Segoe UI", bold=True, size=10)
TOTAL_FILL = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
ALT_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
BRAND_FILL = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
NOTE_FONT = Font(name="Segoe UI", italic=True, size=9, color="64748B")
THIN = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)
USD = '$#,##0.00'
INT = '#,##0'
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")


def _ascii_upper(name: str) -> str:
    decomposed = unicodedata.normalize("NFKD", name)
    stripped = "".join(c for c in decomposed if not unicodedata.combining(c))
    return stripped.upper().replace("İ", "I").replace("ı", "I")


def _find_file(predicate):
    for name in os.listdir(SUPPLIER_DIR):
        if name.startswith("~$"):
            continue
        if predicate(name):
            return SUPPLIER_DIR / name
    raise FileNotFoundError(f"No matching file in {SUPPLIER_DIR}")


def parse_price(raw) -> float:
    if raw is None:
        return 0.0
    if isinstance(raw, (int, float)):
        return float(raw)
    text = str(raw)
    text = text.replace("$", " ").replace("+", " ").replace("VAT", " ").replace("vat", " ")
    text = text.replace(",", ".")
    m = re.search(r"(\d+(?:\.\d+)?)", text)
    if not m:
        raise ValueError(f"Cannot parse price: {raw!r}")
    return float(m.group(1))


def parse_model(raw) -> int:
    if isinstance(raw, (int, float)):
        return int(raw)
    digits = re.sub(r"\D", "", str(raw))
    if not digits:
        raise ValueError(f"Cannot parse model: {raw!r}")
    return int(digits)


def normalize_brand(raw) -> str:
    text = str(raw or "").strip().upper()
    text = text.replace("İ", "I").replace("ı", "I")
    if "BABY" in text:
        return "BABY BASIC"
    if "ARYE" in text or "ARYEH" in text:
        return "ARYE"
    return text or ""


def load_prices(path: Path) -> dict[int, float]:
    wb = load_workbook(path, data_only=True)
    prices: dict[int, float] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            if not row or row[0] is None:
                continue
            key = str(row[0]).strip().upper()
            if key in ("KOD", "CODE", "MODEL"):
                continue
            try:
                model = parse_model(row[0])
            except ValueError:
                continue
            if len(row) < 2 or row[1] is None:
                continue
            try:
                prices[model] = parse_price(row[1])
            except ValueError:
                continue
    if not prices:
        raise RuntimeError(f"No prices found in {path.name}")
    return prices


def load_product_names(path: Path) -> dict[int, str]:
    wb = load_workbook(path, data_only=True)
    names: dict[int, str] = {}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = [str(c).strip().lower() if c is not None else "" for c in rows[0]]
        if "model code" not in header or "product name" not in header:
            continue
        i_name = header.index("product name")
        i_code = header.index("model code")
        for row in rows[1:]:
            if not row or row[i_code] is None:
                continue
            try:
                model = parse_model(row[i_code])
            except ValueError:
                continue
            name = str(row[i_name] or "").strip()
            if name and model not in names:
                names[model] = name
    return names


def load_cartons(path: Path) -> list[dict]:
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    cartons: list[dict] = []
    header_seen = False
    for row in ws.iter_rows(values_only=True):
        if not row or all(c is None or str(c).strip() == "" for c in row[:6]):
            continue
        first = str(row[0]).strip().upper() if row[0] is not None else ""
        if not header_seen:
            if "CARTON" in first or "MODEL" in str(row[1] or "").upper():
                header_seen = True
            continue
        # Subtotal rows have pieces but no model / carton
        if row[1] is None or str(row[1]).strip() == "":
            continue
        try:
            carton = int(float(str(row[0]).strip()))
            model = parse_model(row[1])
            pieces = int(float(str(row[5]).replace(",", "")))
        except (TypeError, ValueError):
            continue
        packets_raw = row[4]
        try:
            packets = int(float(str(packets_raw))) if packets_raw not in (None, "") else 0
        except ValueError:
            packets = 0
        cartons.append({
            "carton": carton,
            "model": model,
            "size": str(row[2] or "").strip(),
            "brand": normalize_brand(row[3]),
            "packets": packets,
            "pieces": pieces,
        })
    if not cartons:
        raise RuntimeError(f"No carton rows found in {path.name}")
    return cartons


def autosize(ws, min_width=10, max_width=36):
    for column in ws.columns:
        max_length = 0
        letter = get_column_letter(column[0].column)
        for cell in column:
            if cell.value is None:
                continue
            max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max(max_length + 2, min_width), max_width)


def style_header_row(ws, row, ncols):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN


def apply_body(ws, start_row, end_row, ncols, money_cols=(), int_cols=()):
    for r in range(start_row, end_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = BODY_FONT
            cell.border = THIN
            cell.alignment = CENTER
            if r % 2 == 0:
                cell.fill = ALT_FILL
            if c in money_cols:
                cell.number_format = USD
            if c in int_cols:
                cell.number_format = INT


def write_summary(wb, by_model, prices, names, total_pcs, total_usd):
    ws = wb.active
    ws.title = "סיכום לתשלום"
    ws.sheet_view.rightToLeft = True
    ws.merge_cells("A1:G1")
    ws["A1"] = "תשלום לספק עומר — משלוח ראשון 07.09.2026"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = RIGHT
    ws.merge_cells("A2:G2")
    ws["A2"] = "מחירים בדולרים ללא מע״מ טורקי · לפי מחירון LEO 2026 ומסמך 200-201-202"
    ws["A2"].font = SUB_FONT
    ws["A2"].alignment = RIGHT

    headers = ["מודל", "מוצר", "יח' Baby Basic", "יח' Arye", "סה״כ יח'", "מחיר $ ליחידה", "סה״כ $"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=4, column=col, value=h)
    style_header_row(ws, 4, len(headers))

    models = sorted(by_model)
    row = 5
    for model in models:
        bb = by_model[model]["BABY BASIC"]
        arye = by_model[model]["ARYE"]
        qty = bb + arye
        price = prices[model]
        ws.cell(row=row, column=1, value=model)
        ws.cell(row=row, column=2, value=names.get(model, ""))
        ws.cell(row=row, column=3, value=bb)
        ws.cell(row=row, column=4, value=arye)
        ws.cell(row=row, column=5, value=qty)
        ws.cell(row=row, column=6, value=price)
        ws.cell(row=row, column=7, value=round(qty * price, 2))
        row += 1
    last_data = row - 1
    apply_body(ws, 5, last_data, 7, money_cols=(6, 7), int_cols=(1, 3, 4, 5))
    ws.cell(row=5, column=2).alignment = RIGHT
    for r in range(5, last_data + 1):
        ws.cell(row=r, column=2).alignment = RIGHT

    for col in range(1, 8):
        cell = ws.cell(row=row, column=col)
        cell.fill = TOTAL_FILL
        cell.font = BOLD_FONT
        cell.border = THIN
        cell.alignment = CENTER
    ws.cell(row=row, column=1, value="סה״כ")
    ws.cell(row=row, column=3, value=sum(by_model[m]["BABY BASIC"] for m in models))
    ws.cell(row=row, column=4, value=sum(by_model[m]["ARYE"] for m in models))
    ws.cell(row=row, column=5, value=total_pcs)
    ws.cell(row=row, column=7, value=round(total_usd, 2))
    for c in (3, 4, 5):
        ws.cell(row=row, column=c).number_format = INT
    ws.cell(row=row, column=7).number_format = USD

    note_row = row + 2
    ws.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=7)
    ws.cell(
        row=note_row,
        column=1,
        value="הערה: מקדמה של $19,314 הועברה דרך חברת השילוח (26.8–3.9.2026) ואינה מקוזזת בקובץ זה.",
    )
    ws.cell(row=note_row, column=1).font = NOTE_FONT
    ws.cell(row=note_row, column=1).alignment = RIGHT
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[4].height = 20
    autosize(ws)


def write_by_brand(wb, by_brand, prices, names):
    ws = wb.create_sheet("לפי מותג")
    ws.sheet_view.rightToLeft = True
    ws.merge_cells("A1:F1")
    ws["A1"] = "פירוט לפי מותג — ללא מע״מ טורקי"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = RIGHT

    headers = ["מותג", "מודל", "מוצר", "סה״כ יח'", "מחיר $ ליחידה", "סה״כ $"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=3, column=col, value=h)
    style_header_row(ws, 3, len(headers))

    row = 4
    brand_order = ["BABY BASIC", "ARYE"]
    brand_totals = {}
    for brand in brand_order:
        models = sorted(by_brand[brand])
        start = row
        brand_pcs = 0
        brand_usd = 0.0
        for model in models:
            qty = by_brand[brand][model]
            price = prices[model]
            amount = round(qty * price, 2)
            ws.cell(row=row, column=1, value=brand)
            ws.cell(row=row, column=2, value=model)
            ws.cell(row=row, column=3, value=names.get(model, ""))
            ws.cell(row=row, column=4, value=qty)
            ws.cell(row=row, column=5, value=price)
            ws.cell(row=row, column=6, value=amount)
            row += 1
            brand_pcs += qty
            brand_usd += amount
        apply_body(ws, start, row - 1, 6, money_cols=(5, 6), int_cols=(2, 4))
        for r in range(start, row):
            ws.cell(row=r, column=3).alignment = RIGHT
        for col in range(1, 7):
            cell = ws.cell(row=row, column=col)
            cell.fill = BRAND_FILL
            cell.font = BOLD_FONT
            cell.border = THIN
            cell.alignment = CENTER
        ws.cell(row=row, column=1, value=f"סה״כ {brand}")
        ws.cell(row=row, column=4, value=brand_pcs)
        ws.cell(row=row, column=4).number_format = INT
        ws.cell(row=row, column=6, value=round(brand_usd, 2))
        ws.cell(row=row, column=6).number_format = USD
        brand_totals[brand] = (brand_pcs, brand_usd)
        row += 2

    grand_row = row
    for col in range(1, 7):
        cell = ws.cell(row=grand_row, column=col)
        cell.fill = TOTAL_FILL
        cell.font = BOLD_FONT
        cell.border = THIN
        cell.alignment = CENTER
    ws.cell(row=grand_row, column=1, value="סה״כ כללי")
    ws.cell(row=grand_row, column=4, value=sum(v[0] for v in brand_totals.values()))
    ws.cell(row=grand_row, column=4).number_format = INT
    ws.cell(row=grand_row, column=6, value=round(sum(v[1] for v in brand_totals.values()), 2))
    ws.cell(row=grand_row, column=6).number_format = USD
    autosize(ws)


def write_cartons(wb, cartons, prices, names):
    ws = wb.create_sheet("פירוט קרטונים")
    ws.sheet_view.rightToLeft = True
    ws.merge_cells("A1:H1")
    ws["A1"] = "53 קרטונים — משלוח ראשון עומר 07.09.2026"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = RIGHT

    headers = ["קרטון", "מודל", "מוצר", "מידה", "מותג", "חבילות", "יחידות", "מחיר $ ליחידה", "סה״כ $"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=3, column=col, value=h)
    style_header_row(ws, 3, len(headers))

    row = 4
    for item in cartons:
        price = prices[item["model"]]
        amount = round(item["pieces"] * price, 2)
        ws.cell(row=row, column=1, value=item["carton"])
        ws.cell(row=row, column=2, value=item["model"])
        ws.cell(row=row, column=3, value=names.get(item["model"], ""))
        ws.cell(row=row, column=4, value=item["size"])
        ws.cell(row=row, column=5, value=item["brand"])
        ws.cell(row=row, column=6, value=item["packets"])
        ws.cell(row=row, column=7, value=item["pieces"])
        ws.cell(row=row, column=8, value=price)
        ws.cell(row=row, column=9, value=amount)
        row += 1
    last = row - 1
    apply_body(ws, 4, last, 9, money_cols=(8, 9), int_cols=(1, 2, 6, 7))
    for r in range(4, last + 1):
        ws.cell(row=r, column=3).alignment = RIGHT

    for col in range(1, 10):
        cell = ws.cell(row=row, column=col)
        cell.fill = TOTAL_FILL
        cell.font = BOLD_FONT
        cell.border = THIN
        cell.alignment = CENTER
    ws.cell(row=row, column=1, value="סה״כ")
    ws.cell(row=row, column=6, value=sum(c["packets"] for c in cartons))
    ws.cell(row=row, column=7, value=sum(c["pieces"] for c in cartons))
    ws.cell(row=row, column=9, value=round(sum(c["pieces"] * prices[c["model"]] for c in cartons), 2))
    ws.cell(row=row, column=6).number_format = INT
    ws.cell(row=row, column=7).number_format = INT
    ws.cell(row=row, column=9).number_format = USD
    autosize(ws, min_width=8)


def register_document():
    with open(SUPPLIERS_JSON, "r", encoding="utf-8") as f:
        suppliers = json.load(f)
    omer = next(s for s in suppliers if s.get("id") == 4)
    docs = omer.setdefault("documents", [])
    if any(d.get("filename") == OUT_NAME for d in docs):
        return False
    next_id = max((d.get("id", 0) for d in docs), default=0) + 1
    docs.append({
        "id": next_id,
        "filename": OUT_NAME,
        "original_name": OUT_NAME,
        "uploaded_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    })
    with open(SUPPLIERS_JSON, "w", encoding="utf-8") as f:
        json.dump(suppliers, f, ensure_ascii=False, indent=2)
        f.write("\n")
    return True


def main():
    price_path = _find_file(lambda n: "PRICE" in _ascii_upper(n) and "LIST" in _ascii_upper(n))
    packing_path = _find_file(lambda n: n.lower() == "200-201-202.xlsx")
    order_path = _find_file(lambda n: "order for omar" in n.lower())

    prices = load_prices(price_path)
    names = load_product_names(order_path)
    cartons = load_cartons(packing_path)

    missing = sorted({c["model"] for c in cartons} - set(prices))
    if missing:
        raise RuntimeError(f"Models in packing list without a price: {missing}")

    by_model = defaultdict(lambda: defaultdict(int))
    by_brand = defaultdict(lambda: defaultdict(int))
    for c in cartons:
        by_model[c["model"]][c["brand"]] += c["pieces"]
        by_brand[c["brand"]][c["model"]] += c["pieces"]

    total_pcs = sum(c["pieces"] for c in cartons)
    total_usd = sum(c["pieces"] * prices[c["model"]] for c in cartons)

    wb = Workbook()
    write_summary(wb, by_model, prices, names, total_pcs, total_usd)
    write_by_brand(wb, by_brand, prices, names)
    write_cartons(wb, cartons, prices, names)
    wb.save(OUT_PATH)

    registered = register_document()
    summary_lines = [
        f"price_file={price_path.name}",
        f"packing_file={packing_path.name}",
        f"prices={prices}",
        f"cartons={len(cartons)} pcs={total_pcs} usd={total_usd:.2f}",
    ]
    for model in sorted(by_model):
        bb = by_model[model]["BABY BASIC"]
        arye = by_model[model]["ARYE"]
        summary_lines.append(
            f"  model {model}: BB={bb} Arye={arye} total={bb+arye} x {prices[model]} = {(bb+arye)*prices[model]:.2f}"
        )
    summary_lines.append(f"wrote={OUT_PATH}")
    summary_lines.append(f"registered={registered}")
    log_path = ROOT / "scripts" / "_omer_payment_build.log"
    log_path.write_text("\n".join(summary_lines), encoding="utf-8")
    print(f"pcs={total_pcs} usd={total_usd:.2f} cartons={len(cartons)} registered={registered}")


if __name__ == "__main__":
    main()
